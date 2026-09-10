"""
Repair the Getty Videos / Getty Stills split of a freshly sorted workbook
using a previous, hand-checked version of the same episode as the truth.

Why this exists: classify.py can only tell a Getty still from a Getty
video by the filename extension or by the EDL's own "Source" value
("Getty Images - Stills"). Some EDL exports carry neither -- every Getty
row says "Getty Images - Footage" and the stills come through as bare ids
("GETTYIMAGES-113598072", "GETTYIMAGES-2217807802-.NEW.01") with no .jpg
-- so those stills land on the Getty Videos sheet. The old delivery of the
same episode was corrected by hand, and its "Getty pics" sheet is the only
record of which ids are actually stills.

So: for every clip on the new file's two Getty sheets, if the old file
files that clip under the OTHER one of the two, the row is moved. Clips
the old file doesn't have at all (new material) are never moved -- there
is nothing to check them against -- and are returned under "unmatched" so
they can be eyeballed.

Deliberately narrow, don't widen it:
  - Only "Getty Videos" and "Getty Stills"/"Getty pics" are touched. Every
    other sheet is copied through verbatim, including "Worksheet".
  - Only the row's SHEET changes. No cell value is rewritten -- notably not
    "Source"/"Source Reel Name", which is EDL data ("Getty Images -
    Footage" even on the misfiled stills, which is exactly why they were
    misfiled). Only the zebra fill is recomputed, from each row's new
    position.
  - Rows keep their original relative order within a sheet.

Matching is dedupe.dedup_key(), case-insensitive: old masters uppercase
the ids the new EDL writes in mixed case, and only one of the two carries
an extension.

Input is a SORTED workbook (the output of `bigmedia sort`), not a grouped
one -- run `bigmedia group` on the result afterwards (the `fix-getty` CLI
command does that for you) so the Total Duration/Seconds columns and the
totals block are recomputed for both Getty sheets.
"""
import copy
import re
from openpyxl import load_workbook, Workbook

from .dedupe import dedup_key
from .xlsx_utils import find_column_any, copy_sheet_verbatim, style_output_sheets, first_data_font

NAME_COLUMN_ALIASES = ("Clip Name", "Name")

VIDEOS_TITLE = "Getty Videos"
STILLS_TITLE = "Getty Stills"
VIDEO_SHEET_ALIASES = ("getty videos",)
# "Getty pics" is what the older hand-made masters call the stills sheet.
STILLS_SHEET_ALIASES = ("getty stills", "getty pics")

_EPISODE_RE = re.compile(r'EP\s*0*(\d+)', re.IGNORECASE)


def _find_sheet(wb, aliases):
    for title in wb.sheetnames:
        if title.strip().lower() in aliases:
            return wb[title]
    return None


def _name_col(ws, name_column):
    candidates = [name_column] + [a for a in NAME_COLUMN_ALIASES if a != name_column]
    return find_column_any(ws, candidates)


def _clip_rows(ws, name_col_idx):
    """(row, name) for every row holding a clip. The trailing "Total
    clips"/"Total Seconds" block a grouped workbook leaves in the name
    column is skipped, so this is safe to point at either shape."""
    for r in range(2, ws.max_row + 1):
        value = ws.cell(row=r, column=name_col_idx).value
        if value is None or str(value).strip() == "":
            continue
        if str(value).strip().lower().startswith("total "):
            continue
        yield r, value


def _keys(ws, name_column):
    """{dedup key: first filename seen} for a sheet, or {} if absent."""
    if ws is None:
        return {}
    try:
        name_col_idx = _name_col(ws, name_column)
    except ValueError:
        return {}
    keys = {}
    for _, value in _clip_rows(ws, name_col_idx):
        keys.setdefault(dedup_key(str(value)).casefold(), str(value))
    return keys


def episode_number(path, wb=None):
    """Episode number of a workbook: the "Episode 20" value the sorted
    files carry in their first column, else an "EP20" in the filename.
    None if neither is present."""
    if wb is not None:
        for ws in wb.worksheets:
            if str(ws.cell(row=1, column=1).value or "").strip().lower() != "episode":
                continue
            for r in range(2, min(ws.max_row, 20) + 1):
                match = re.search(r'(\d+)', str(ws.cell(row=r, column=1).value or ""))
                if match:
                    return int(match.group(1))
    match = _EPISODE_RE.search(str(path))
    return int(match.group(1)) if match else None


def pair_by_episode(old_paths, new_paths):
    """[(episode, old_path, new_path)] for episodes present in both
    folders, plus the leftovers. Filenames are useless for pairing here
    ("EP4 - Submarine (kopie 2).xlsx" vs "SECRETS OF - MASTER EDL -
    SUFFREN NUCLEAR SUBMARINE_sorted.xlsx"), the episode number is not."""
    def index(paths):
        found = {}
        for path in paths:
            wb = load_workbook(path, read_only=True)
            try:
                episode = episode_number(path, wb)
            finally:
                wb.close()
            if episode is not None:
                found.setdefault(episode, path)
        return found

    old_index, new_index = index(old_paths), index(new_paths)
    pairs = [(ep, old_index[ep], new_index[ep]) for ep in sorted(set(old_index) & set(new_index))]
    unpaired_old = [old_index[ep] for ep in sorted(set(old_index) - set(new_index))]
    unpaired_new = [new_index[ep] for ep in sorted(set(new_index) - set(old_index))]
    return pairs, unpaired_old, unpaired_new


def _copy_cell(src_cell, dst_cell):
    dst_cell.value = src_cell.value
    if src_cell.has_style:
        dst_cell.font = copy.copy(src_cell.font)
        dst_cell.border = copy.copy(src_cell.border)
        dst_cell.alignment = copy.copy(src_cell.alignment)
        dst_cell.number_format = src_cell.number_format
        dst_cell.protection = copy.copy(src_cell.protection)


def _zebra_fills(ws):
    """The two alternating row fills of a sorted sheet, sampled from its
    first two data rows so a rewritten sheet keeps the same striping."""
    return copy.copy(ws.cell(row=2, column=1).fill), copy.copy(ws.cell(row=3, column=1).fill)


def _write_sheet(wb_out, title, header_ws, rows, fills):
    """Write one Getty sheet: header cloned from header_ws, then `rows`
    ((source worksheet, row number) pairs) in order, re-striped."""
    ws_out = wb_out.create_sheet(title=title)
    max_col = header_ws.max_column
    for c in range(1, max_col + 1):
        src = header_ws.cell(row=1, column=c)
        dst = ws_out.cell(row=1, column=c)
        _copy_cell(src, dst)
        dst.fill = copy.copy(src.fill)
    if 1 in header_ws.row_dimensions:
        ws_out.row_dimensions[1].height = header_ws.row_dimensions[1].height

    fill_even, fill_odd = fills
    for out_r, (ws_src, src_r) in enumerate(rows, start=2):
        fill = fill_even if out_r % 2 == 0 else fill_odd
        for c in range(1, max_col + 1):
            dst = ws_out.cell(row=out_r, column=c)
            _copy_cell(ws_src.cell(row=src_r, column=c), dst)
            dst.fill = copy.copy(fill)
        if src_r in ws_src.row_dimensions:
            ws_out.row_dimensions[out_r].height = ws_src.row_dimensions[src_r].height

    for col_letter, dim in header_ws.column_dimensions.items():
        ws_out.column_dimensions[col_letter].width = dim.width
    ws_out.freeze_panes = "A2"
    return ws_out


def fix_getty_split(old_path, new_path, out_path, name_column="Clip Name"):
    """Rewrite `new_path` to `out_path` with its Getty Videos/Getty Stills
    rows re-filed according to `old_path`.

    Returns {"video_to_stills", "stills_to_video": [clip names moved],
    "rows_moved_to_stills", "rows_moved_to_video": row counts,
    "unmatched_videos", "unmatched_stills": [clip names the old file has on
    neither Getty sheet, left where they are],
    "stills_sheet_created": bool}.
    """
    wb_old = load_workbook(old_path, data_only=True)
    wb_new = load_workbook(new_path)

    old_video_keys = set(_keys(_find_sheet(wb_old, VIDEO_SHEET_ALIASES), name_column))
    old_stills_keys = set(_keys(_find_sheet(wb_old, STILLS_SHEET_ALIASES), name_column))
    wb_old.close()

    ws_videos = _find_sheet(wb_new, VIDEO_SHEET_ALIASES)
    ws_stills = _find_sheet(wb_new, STILLS_SHEET_ALIASES)
    if ws_videos is None and ws_stills is None:
        raise ValueError(f"{new_path}: no Getty Videos or Getty Stills sheet to fix")

    video_rows, stills_rows = [], []
    result = {
        "video_to_stills": [], "stills_to_video": [],
        "rows_moved_to_stills": 0, "rows_moved_to_video": 0,
        "unmatched_videos": [], "unmatched_stills": [],
        "stills_sheet_created": False, "videos_sheet_created": False,
    }

    def triage(ws, own_bucket, other_bucket, old_other_keys, old_own_keys,
               moved_names, moved_rows_key, unmatched_names):
        if ws is None:
            return
        name_col_idx = _name_col(ws, name_column)
        for r, value in _clip_rows(ws, name_col_idx):
            key = dedup_key(str(value)).casefold()
            if key in old_other_keys:
                other_bucket.append((ws, r))
                result[moved_rows_key] += 1
                if str(value) not in moved_names:
                    moved_names.append(str(value))
            else:
                own_bucket.append((ws, r))
                if key not in old_own_keys and str(value) not in unmatched_names:
                    unmatched_names.append(str(value))

    triage(ws_videos, video_rows, stills_rows, old_stills_keys, old_video_keys,
           result["video_to_stills"], "rows_moved_to_stills", result["unmatched_videos"])
    triage(ws_stills, stills_rows, video_rows, old_video_keys, old_stills_keys,
           result["stills_to_video"], "rows_moved_to_video", result["unmatched_stills"])

    # Each sheet keeps its own rows first, in their original order, with the
    # rows moved in from the other sheet appended after them (also in their
    # original order) -- so a corrected file reads as "what was already
    # right, then what was brought over".
    video_rows.sort(key=lambda item: (item[0] is not ws_videos, item[1]))
    stills_rows.sort(key=lambda item: (item[0] is not ws_stills, item[1]))

    videos_title = ws_videos.title if ws_videos is not None else VIDEOS_TITLE
    stills_title = ws_stills.title if ws_stills is not None else STILLS_TITLE
    # An EDL whose stills all came through extension-less has no stills
    # sheet at all; build one from the other Getty sheet's header/styling.
    # (Same the other way round, for a file with stills but no videos.)
    if ws_stills is None and stills_rows:
        result["stills_sheet_created"] = True
    if ws_videos is None and video_rows:
        result["videos_sheet_created"] = True

    wb_out = Workbook()
    wb_out.remove(wb_out.active)
    for title in wb_new.sheetnames:
        ws_src = wb_new[title]
        if ws_src is ws_videos:
            _write_sheet(wb_out, videos_title, ws_videos, video_rows, _zebra_fills(ws_videos))
            if result["stills_sheet_created"]:
                _write_sheet(wb_out, stills_title, ws_videos, stills_rows, _zebra_fills(ws_videos))
        elif ws_src is ws_stills:
            if result.get("videos_sheet_created"):
                _write_sheet(wb_out, videos_title, ws_stills, video_rows, _zebra_fills(ws_stills))
            _write_sheet(wb_out, stills_title, ws_stills, stills_rows, _zebra_fills(ws_stills))
        else:
            copy_sheet_verbatim(ws_src, wb_out.create_sheet(title=title), ws_src.max_row, ws_src.max_column)

    getty_sheets = [wb_out[t] for t in wb_out.sheetnames
                    if t in (videos_title, stills_title)]
    if getty_sheets:
        # index-keyed widths assume both Getty sheets share a column layout
        # -- true for `sort` output, the documented input to this command.
        style_output_sheets(getty_sheets, width_by="index",
                            data_font=first_data_font(getty_sheets))

    wb_out.save(out_path)
    return result
