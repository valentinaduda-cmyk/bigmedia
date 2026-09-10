"""
Shared openpyxl helpers used by every workbook-processing command.

Keeping these here means sort_workbook.py, dedupe.py, and any future
command (report generation, QC checks, ...) don't each reinvent "copy this
sheet with its styling" from scratch.
"""
import copy
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill


def iter_xlsx_files(path, exclude_suffix=None):
    """Path objects for the .xlsx file(s) at `path`. If `path` is a folder,
    every .xlsx file directly inside it is included, except Excel's own
    "~$foo.xlsx" lock files for currently-open workbooks. If `exclude_suffix`
    is given (e.g. "grouped"), files whose stem ends with "_<exclude_suffix>"
    are also skipped -- pass the current command's own output suffix so a
    batch command re-run on the same folder doesn't reprocess its own prior
    output every time (which would otherwise snowball file count on each
    run: "x.xlsx" -> "x_grouped.xlsx" -> "x_grouped_grouped.xlsx" -> ...).
    Other stages' suffixes (e.g. "_sorted" files being fed into "group") are
    legitimate input and must NOT be filtered out. If `path` is a single
    file, it's returned as a one-item list unfiltered."""
    p = Path(path)
    if p.is_dir():
        files = (f for f in p.glob("*.xlsx") if not f.name.startswith("~$"))
        if exclude_suffix:
            files = (f for f in files if not f.stem.endswith(f"_{exclude_suffix}"))
        return sorted(files)
    return [p]


def find_column(ws, header_name: str) -> int:
    """1-indexed column number of the first header cell matching header_name
    (case-sensitive, exact match) in row 1. Raises ValueError if not found."""
    for c in range(1, ws.max_column + 1):
        if ws.cell(row=1, column=c).value == header_name:
            return c
    raise ValueError(f"Column {header_name!r} not found in header row")


def find_column_any(ws, header_names) -> int:
    """Like find_column, but tries each name in header_names in order and
    returns the first match. Needed because real delivery files aren't
    consistent about column naming (e.g. some use "Name", others "Clip
    Name" for the same data). Raises ValueError if none match."""
    for header_name in header_names:
        for c in range(1, ws.max_column + 1):
            if ws.cell(row=1, column=c).value == header_name:
                return c
    raise ValueError(f"None of {list(header_names)!r} found in header row")


def capture_header_template(ws):
    """Snapshot everything needed to reproduce this sheet's header row and
    per-column styling elsewhere: header cells, a font/number-format sample
    from row 2, and configured column widths."""
    max_col = ws.max_column
    header_cells = [ws.cell(row=1, column=c) for c in range(1, max_col + 1)]
    col_font = [copy.copy(ws.cell(row=2, column=c).font) for c in range(1, max_col + 1)]
    col_numfmt = [ws.cell(row=2, column=c).number_format for c in range(1, max_col + 1)]
    col_widths = {
        get_column_letter(c): ws.column_dimensions[get_column_letter(c)].width
        for c in range(1, max_col + 1)
        if get_column_letter(c) in ws.column_dimensions
    }
    return {
        "header_cells": header_cells,
        "col_font": col_font,
        "col_numfmt": col_numfmt,
        "col_widths": col_widths,
        "max_col": max_col,
    }


def write_header_row(ws_out, header_cells, row_height=None):
    """Write a styled header row (row 1) onto ws_out from captured header cells."""
    for c, hcell in enumerate(header_cells, start=1):
        nc = ws_out.cell(row=1, column=c, value=hcell.value)
        nc.font = copy.copy(hcell.font)
        nc.fill = copy.copy(hcell.fill)
        nc.alignment = copy.copy(hcell.alignment)
        nc.number_format = hcell.number_format
    if row_height is not None:
        ws_out.row_dimensions[1].height = row_height


def write_data_row(ws_out, out_row, values, col_font, col_numfmt, fill):
    """Write one data row with the captured per-column font/number-format
    and a single fill (e.g. alternating zebra-striping) applied to every cell."""
    for c, val in enumerate(values, start=1):
        nc = ws_out.cell(row=out_row, column=c, value=val)
        nc.font = col_font[c - 1]
        nc.number_format = col_numfmt[c - 1]
        nc.fill = fill


def copy_sheet_verbatim(src_ws, dst_ws, max_row=None, max_col=None):
    """Clone src_ws into dst_ws exactly as-is: values, per-cell styling, row
    heights, column widths, merged cells, and freeze panes. Used for
    untouched backup copies of raw input data."""
    max_row = max_row or src_ws.max_row
    max_col = max_col or src_ws.max_column

    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            src_cell = src_ws.cell(row=r, column=c)
            dst_cell = dst_ws.cell(row=r, column=c, value=src_cell.value)
            if src_cell.has_style:
                dst_cell.font = copy.copy(src_cell.font)
                dst_cell.fill = copy.copy(src_cell.fill)
                dst_cell.border = copy.copy(src_cell.border)
                dst_cell.alignment = copy.copy(src_cell.alignment)
                dst_cell.number_format = src_cell.number_format
                dst_cell.protection = copy.copy(src_cell.protection)
        if r in src_ws.row_dimensions:
            dst_ws.row_dimensions[r].height = src_ws.row_dimensions[r].height

    for col_letter, dim in src_ws.column_dimensions.items():
        dst_ws.column_dimensions[col_letter].width = dim.width

    for merged_range in src_ws.merged_cells.ranges:
        dst_ws.merge_cells(str(merged_range))

    dst_ws.freeze_panes = src_ws.freeze_panes


# Excel column widths are measured in characters of the default font, so a
# character count plus a little padding is a good enough proxy for "wide
# enough to read without dragging the edge".
AUTOFIT_PADDING = 2
AUTOFIT_MIN_WIDTH = 10
AUTOFIT_MAX_WIDTH = 60


def measure_column_widths(rows, headers=None):
    """Width per column index (1-based) sized to the longest value in it.

    `rows` is an iterable of value sequences; `headers` is an optional extra
    sequence measured alongside them, so a column whose header is longer than
    any of its data still fits. Clamped to [AUTOFIT_MIN_WIDTH, AUTOFIT_MAX_WIDTH]:
    without a max, one pathological clip name makes the sheet unusable sideways.
    """
    longest = {}
    sources = []
    if headers:
        sources.append(headers)
    sources.extend(rows)
    for values in sources:
        for c, val in enumerate(values, start=1):
            if val is None:
                continue
            n = max(len(part) for part in str(val).split("\n"))
            if n > longest.get(c, 0):
                longest[c] = n
    return {
        c: min(max(n + AUTOFIT_PADDING, AUTOFIT_MIN_WIDTH), AUTOFIT_MAX_WIDTH)
        for c, n in longest.items()
    }


def apply_column_widths(ws, widths):
    """Set column widths from a {column index: width} mapping."""
    for c, width in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = width


def apply_uniform_data_font(ws, font, min_row=2):
    """Give every data cell on the sheet the same font, leaving the header
    row alone. Source workbooks routinely mix fonts column to column, which
    makes the output read as several tables stapled together."""
    for row in ws.iter_rows(min_row=min_row, max_row=ws.max_row,
                            max_col=ws.max_column):
        for cell in row:
            cell.font = copy.copy(font)


def measure_header_widths(worksheets):
    """Width per HEADER NAME, measured across several sheets at once.

    Keyed by header rather than column index because commands that insert
    columns mid-table (group_duplicates) leave the same header sitting at
    different indices on different sheets -- keying by index would make
    "Clip Name" a different width depending on which tab you opened.
    """
    longest = {}
    for ws in worksheets:
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        for c, header in enumerate(headers, start=1):
            if header is None:
                continue
            best = longest.get(header, len(str(header)))
            for r in range(1, ws.max_row + 1):
                val = ws.cell(row=r, column=c).value
                if val is None:
                    continue
                n = max(len(part) for part in str(val).split("\n"))
                if n > best:
                    best = n
            longest[header] = best
    return {
        h: min(max(n + AUTOFIT_PADDING, AUTOFIT_MIN_WIDTH), AUTOFIT_MAX_WIDTH)
        for h, n in longest.items()
    }


def apply_header_widths(ws, widths):
    """Set column widths from a {header name: width} mapping, matching each
    column by the value in its header row."""
    for c in range(1, ws.max_column + 1):
        header = ws.cell(row=1, column=c).value
        if header in widths:
            ws.column_dimensions[get_column_letter(c)].width = widths[header]


# Real delivery headers are white bold text on a dark fill; the styling is
# now unconditional -- a fixed black fill with a white bold font.
HEADER_FILL = PatternFill("solid", fgColor="FF000000")
_HEADER_FONT_COLOR = "FFFFFFFF"
_FALLBACK_DATA_FONT = Font(name="Calibri", size=11)


def apply_header_style(ws):
    """Row 1: every non-empty header cell gets the fixed black fill and a
    white bold font, keeping the cell's own family and size."""
    for c in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=c)
        if cell.value is None:
            continue
        cell.fill = copy.copy(HEADER_FILL)
        f = cell.font
        cell.font = Font(name=f.name, size=f.size, bold=True, color=_HEADER_FONT_COLOR)


def sample_data_font(ws, col_idx):
    """The row-2 font of column `col_idx`, to standardize data cells on.
    Calibri 11 when there is no row 2 or the cell carries no font."""
    if ws.max_row < 2:
        return copy.copy(_FALLBACK_DATA_FONT)
    f = ws.cell(row=2, column=col_idx).font
    if f is None or f.name is None:
        return copy.copy(_FALLBACK_DATA_FONT)
    return copy.copy(f)


def first_data_font(worksheets):
    """First non-empty data cell (row 2+) font across the given sheets;
    Calibri 11 fallback."""
    for ws in worksheets:
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
            for cell in row:
                if cell.value is not None and cell.font is not None and cell.font.name is not None:
                    return copy.copy(cell.font)
    return copy.copy(_FALLBACK_DATA_FONT)


def style_output_sheets(sheets, *, width_by, data_font):
    """The single styling pass every in-scope command runs on its finished
    output sheets. Skips any sheet titled "Worksheet". `width_by` is
    "index" (columns fixed across sheets) or "header" (columns shift, key
    widths by header name). Applies shared column widths clamped to
    [10, 60], the uniform `data_font` to row 2+, and apply_header_style to
    row 1."""
    targets = [ws for ws in sheets if ws.title != "Worksheet"]
    if not targets:
        return
    if width_by == "index":
        # Feed every row -- header row included -- of every target sheet as a
        # plain row. Header text and data are then positionally self-aligned
        # within each sheet, and the cross-sheet max is still taken. (Using
        # the `headers=` param with a flat concat of all sheets' headers would
        # misalign column indices when sibling sheets differ in column count.)
        widths = measure_column_widths(
            [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
            for ws in targets
            for r in range(1, ws.max_row + 1)
        )
        for ws in targets:
            apply_column_widths(ws, widths)
    elif width_by == "header":
        widths = measure_header_widths(targets)
        for ws in targets:
            apply_header_widths(ws, widths)
    else:
        raise ValueError(f"width_by must be 'index' or 'header', got {width_by!r}")
    for ws in targets:
        apply_uniform_data_font(ws, data_font)
        apply_header_style(ws)


def analyze_files(paths):
    """Read-only inspection of one or more clip-list workbooks, for the web
    UI's analyze-and-suggest step. Returns the union of the active sheets'
    header names, every sheet's name and (header-excluded) row count summed
    across files, and warnings for files that won't open or whose columns
    disagree with the first readable file. Never raises for a bad file --
    it lands in "warnings" and the others are still processed.

    Header values are stringified (``str(value).strip()``) so a stray
    ``datetime``/number cell in row 1 still yields a JSON-serializable
    header; whitespace-only cells are still dropped."""
    headers = []
    seen_headers = set()
    first_header_set = None
    first_name = None
    sheet_rows = {}
    sheet_order = []
    warnings = []

    for path in paths:
        name = Path(path).name
        wb = None
        try:
            wb = load_workbook(path, read_only=True, data_only=True)

            active = wb.active
            file_headers = [
                str(c.value).strip() for c in next(active.iter_rows(min_row=1, max_row=1), [])
                if c.value is not None and str(c.value).strip() != ""
            ]
            for h in file_headers:
                if h not in seen_headers:
                    seen_headers.add(h)
                    headers.append(h)

            if first_header_set is None:
                first_header_set = set(file_headers)
                first_name = name
            elif set(file_headers) != first_header_set:
                warnings.append(f"{name}: columns differ from {first_name}")

            for title in wb.sheetnames:
                ws = wb[title]
                rows = ws.max_row
                if rows is None:
                    rows = sum(1 for _ in ws.iter_rows())
                count = max(rows - 1, 0)
                if title not in sheet_rows:
                    sheet_rows[title] = 0
                    sheet_order.append(title)
                sheet_rows[title] += count

        except Exception as exc:  # openpyxl raises several unrelated types
            warnings.append(f"{name}: could not read ({exc})")
        finally:
            if wb is not None:
                wb.close()

    return {
        "headers": headers,
        "sheets": [{"name": t, "rows": sheet_rows[t]} for t in sheet_order],
        "warnings": warnings,
    }
