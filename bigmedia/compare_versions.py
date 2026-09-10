"""
Version diff between two deliveries of the same episode, per source
category: how many clips each category holds on either side, which clips
the new version gained, and which it lost.

The team re-cuts an episode every few weeks and re-runs `bigmedia sort` on
it. After each round the questions are always the same, per category:
is the count still what I expect, what came in, what went away.

Four buckets per category, and the distinction between them is the whole
point -- a clip that just changed category is not new material and must
not be chased twice:

    added      key is on this category of the new file and on NO sheet of
               the old file  -> genuinely new material
    moved in   key is on this category of the new file and somewhere else
               on the old file -> recategorized, already cleared
    removed    key is on this category of the old file and on NO sheet of
               the new file -> dropped from the cut
    moved out  key is on this category of the old file and somewhere else
               on the new file -> recategorized, still in the cut

Matching is dedupe.dedup_key() (extension, "(1)" copy markers and export
junk stripped), case-insensitive: old masters uppercase the ids the new
EDL writes in mixed case. Pass case_sensitive=True to compare verbatim.

Categories are paired by sheet name, case-insensitively, honouring the
"Getty pics"/"Getty Stills" and "GFX"/"graphics" aliases -- the two files
being compared are often months apart and didn't name their sheets
identically. A category present on only one side is still reported, with
zero on the other.

Backup tabs ("Worksheet", "... Master XML", "Kopie listu ...") are not
categories and get no row of their own, but they DO count as "the clip is
somewhere in this file" when deciding added vs moved -- that is exactly
what they are evidence of.

Output: a "Summary" sheet, one row per category with both sides' counts
and the four bucket counts, followed by "<category> added" and
"<category> removed" sheets carrying the actual rows (from the new and the
old file respectively, with that file's own columns and styling) behind a
"Status" column that says which bucket the row is in.
"""
import copy
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter, column_index_from_string

from .dedupe import dedup_key
from .xlsx_utils import (
    find_column_any, capture_header_template, write_data_row,
    style_output_sheets, apply_header_style, sample_data_font,
)

# Same header aliases the other commands accept -- real files use either.
NAME_COLUMN_ALIASES = ("Clip Name", "Name")

# Sheets that are a copy of the whole EDL rather than a source category:
# the backup tab `sort`/`dedupe` write, and the raw master/XML tabs the
# edit team leaves in hand-made workbooks (Czech "Kopie listu" = "copy of
# sheet").
BACKUP_SHEET_PATTERNS = ("worksheet", "master xml", "kopie listu", "copy of")

# Categories that mean the same thing under different names across
# versions, so they pair up instead of showing as one all-removed and one
# all-added category.
SHEET_ALIASES = {
    "getty pics": "getty stills",
    "getty stills": "getty stills",
    "getty videos": "getty videos",
    "graphics": "gfx",
    "gfx": "gfx",
}

SUMMARY_TITLE = "Summary"
SUMMARY_HEADERS = [
    "Category", "Clips in old", "Clips in new", "Added", "Removed",
    "Moved in (from another old sheet)", "Moved out (to another new sheet)",
]
STATUS_HEADER = "Status"
MAX_SHEET_TITLE = 31


def is_backup_sheet(title) -> bool:
    t = (title or "").strip().lower()
    return any(p in t for p in BACKUP_SHEET_PATTERNS)


def category_key(title) -> str:
    """The name a sheet is paired on across versions."""
    t = (title or "").strip().lower()
    return SHEET_ALIASES.get(t, t)


def compare_key(name, case_sensitive=False) -> str:
    """The identity of a clip for diffing purposes. See module docstring."""
    key = dedup_key(str(name))
    return key if case_sensitive else key.casefold()


def _name_col(ws, name_column):
    """Index of the clip-name column, or None if this sheet has no header
    for one (chart/notes tabs people leave in a workbook)."""
    candidates = [name_column] + [a for a in NAME_COLUMN_ALIASES if a != name_column]
    try:
        return find_column_any(ws, candidates)
    except ValueError:
        return None


def _clip_rows(ws, name_col_idx):
    """(row, name) for every row holding a clip. Skips the "Total clips" /
    "Total Seconds" block a grouped workbook leaves in the name column, so
    this reads sorted and grouped files alike."""
    for r in range(2, ws.max_row + 1):
        value = ws.cell(row=r, column=name_col_idx).value
        if value is None or str(value).strip() == "":
            continue
        if str(value).strip().lower().startswith("total "):
            continue
        yield r, value


def _scan(wb, name_column, case_sensitive):
    """(categories, all_keys) for one workbook.

    categories maps category key -> {"title", "sheet", "name_col",
    "rows": {clip key: [row numbers]}}, backup tabs excluded; all_keys is
    every key in the file, backup tabs included.
    """
    categories = {}
    all_keys = set()
    for ws in wb.worksheets:
        name_col_idx = _name_col(ws, name_column)
        if name_col_idx is None:
            continue
        rows = {}
        for r, value in _clip_rows(ws, name_col_idx):
            key = compare_key(value, case_sensitive)
            all_keys.add(key)
            rows.setdefault(key, []).append(r)
        if is_backup_sheet(ws.title):
            continue
        cat = category_key(ws.title)
        if cat in categories:
            # Two sheets folding to one category (e.g. "Getty pics" and
            # "Getty Stills" side by side): merge, first sheet keeps the title.
            for key, row_numbers in rows.items():
                categories[cat]["rows"].setdefault(key, []).extend(row_numbers)
            continue
        categories[cat] = {"title": ws.title, "sheet": ws, "name_col": name_col_idx, "rows": rows}
    return categories, all_keys


def _where(categories, key, exclude_cat):
    """Title of another category holding this key, for the Status text."""
    for cat, data in categories.items():
        if cat != exclude_cat and key in data["rows"]:
            return data["title"]
    return "?"


def _write_summary(wb_out, report):
    ws = wb_out.create_sheet(title=SUMMARY_TITLE)
    for c, header in enumerate(SUMMARY_HEADERS, start=1):
        ws.cell(row=1, column=c, value=header).font = Font(bold=True)
    for r, (title, counts) in enumerate(report.items(), start=2):
        for c, value in enumerate([
            title, counts["old_clips"], counts["new_clips"],
            len(counts["added"]), len(counts["removed"]),
            len(counts["moved_in"]), len(counts["moved_out"]),
        ], start=1):
            ws.cell(row=r, column=c, value=value)
    total_row = len(report) + 3
    ws.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True)
    for c, field in enumerate(["old_clips", "new_clips", "added", "removed", "moved_in", "moved_out"], start=2):
        value = sum(counts[field] if isinstance(counts[field], int) else len(counts[field])
                    for counts in report.values())
        ws.cell(row=total_row, column=c, value=value).font = Font(bold=True)
    ws.column_dimensions["A"].width = 24
    for c in range(2, len(SUMMARY_HEADERS) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 20
    ws.freeze_panes = "A2"


def _write_rows_sheet(wb_out, title, ws_src, rows):
    """One "<category> added"/"removed" sheet: the source file's own header
    and column styling, shifted one column right behind a Status column."""
    template = capture_header_template(ws_src)
    max_col = template["max_col"]
    ws_out = wb_out.create_sheet(title=title[:MAX_SHEET_TITLE])

    ws_out.cell(row=1, column=1, value=STATUS_HEADER).font = Font(bold=True)
    for c, hcell in enumerate(template["header_cells"], start=2):
        cell = ws_out.cell(row=1, column=c, value=hcell.value)
        cell.font = copy.copy(hcell.font)
        cell.fill = copy.copy(hcell.fill)
        cell.alignment = copy.copy(hcell.alignment)
        cell.number_format = hcell.number_format
    if 1 in ws_src.row_dimensions:
        ws_out.row_dimensions[1].height = ws_src.row_dimensions[1].height

    fill_even = copy.copy(ws_src.cell(row=2, column=1).fill)
    fill_odd = copy.copy(ws_src.cell(row=3, column=1).fill)
    for out_r, (src_r, status) in enumerate(rows, start=2):
        values = [ws_src.cell(row=src_r, column=c).value for c in range(1, max_col + 1)]
        fill = fill_even if out_r % 2 == 0 else fill_odd
        write_data_row(ws_out, out_r, [status] + values,
                       [Font(bold=True)] + template["col_font"],
                       ["General"] + template["col_numfmt"], fill)

    ws_out.column_dimensions["A"].width = 26
    for col_letter, width in template["col_widths"].items():
        shifted = get_column_letter(column_index_from_string(col_letter) + 1)
        ws_out.column_dimensions[shifted].width = width
    ws_out.freeze_panes = "B2"
    return ws_out


def compare_workbooks(old_path, new_path, out_path, name_column="Clip Name",
                      case_sensitive=False, unique=False):
    """Write the per-category old-vs-new report for one episode.

    Returns {category title: {"old_clips", "new_clips" (unique clip counts),
    "old_rows", "new_rows", "added", "removed", "moved_in", "moved_out"
    (lists of clip names)}}, new-file categories first.
    """
    wb_old = load_workbook(old_path, data_only=True)
    wb_new = load_workbook(new_path, data_only=True)
    old_cats, old_all = _scan(wb_old, name_column, case_sensitive)
    new_cats, new_all = _scan(wb_new, name_column, case_sensitive)

    # New file's categories first, in its own sheet order, then categories
    # only the old file had.
    ordered = list(new_cats) + [c for c in old_cats if c not in new_cats]

    report = {}
    sheets_to_write = []
    for cat in ordered:
        old_data = old_cats.get(cat)
        new_data = new_cats.get(cat)
        old_keys = set(old_data["rows"]) if old_data else set()
        new_keys = set(new_data["rows"]) if new_data else set()
        title = (new_data or old_data)["title"]

        added = sorted(new_keys - old_all)
        moved_in = sorted((new_keys & old_all) - old_keys)
        removed = sorted(old_keys - new_all)
        moved_out = sorted((old_keys & new_all) - new_keys)

        def names(data, keys):
            return [str(data["sheet"].cell(row=data["rows"][k][0], column=data["name_col"]).value)
                    for k in keys]

        report[title] = {
            "old_clips": len(old_keys),
            "new_clips": len(new_keys),
            "old_rows": sum(len(v) for v in old_data["rows"].values()) if old_data else 0,
            "new_rows": sum(len(v) for v in new_data["rows"].values()) if new_data else 0,
            "added": names(new_data, added) if new_data else [],
            "moved_in": names(new_data, moved_in) if new_data else [],
            "removed": names(old_data, removed) if old_data else [],
            "moved_out": names(old_data, moved_out) if old_data else [],
        }

        if new_data and (added or moved_in):
            rows = []
            for key in added + moved_in:
                status = "added" if key in set(added) else f"moved from {_where(old_cats, key, cat)}"
                row_numbers = new_data["rows"][key]
                rows.extend((r, status) for r in (row_numbers[:1] if unique else row_numbers))
            rows.sort(key=lambda item: item[0])
            sheets_to_write.append((f"{title} added", new_data["sheet"], rows))

        if old_data and (removed or moved_out):
            rows = []
            for key in removed + moved_out:
                status = "removed" if key in set(removed) else f"moved to {_where(new_cats, key, cat)}"
                row_numbers = old_data["rows"][key]
                rows.extend((r, status) for r in (row_numbers[:1] if unique else row_numbers))
            rows.sort(key=lambda item: item[0])
            sheets_to_write.append((f"{title} removed", old_data["sheet"], rows))

    wb_out = Workbook()
    wb_out.remove(wb_out.active)
    _write_summary(wb_out, report)
    for title, ws_src, rows in sheets_to_write:
        _write_rows_sheet(wb_out, title, ws_src, rows)

    apply_header_style(wb_out["Summary"])
    row_sheets = [wb_out[t] for t in wb_out.sheetnames if t != "Summary"]
    if row_sheets and sheets_to_write:
        style_output_sheets(
            row_sheets,
            width_by="header",
            data_font=sample_data_font(sheets_to_write[0][1], 1),
        )
    wb_out.save(out_path)
    return report
