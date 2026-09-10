"""Tests for the per-category version-diff command.

The fixtures reproduce the real pair of files this was built for: an old
hand-made master that names its sheets differently ("Getty pics") and
uppercases its filenames, and a new sorted workbook with a "Worksheet"
backup tab. Each case encodes one bucket rule (added / removed / moved in
/ moved out).
"""
import pytest
from openpyxl import Workbook, load_workbook

from bigmedia.compare_versions import compare_workbooks

HEADERS = ["Episode", "Name", "Sequence In", "Clip Duration", "Source"]


def _write(path, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for title, names in sheets:
        ws = wb.create_sheet(title=title)
        for c, h in enumerate(HEADERS, start=1):
            ws.cell(row=1, column=c, value=h)
        for r, name in enumerate(names, start=2):
            ws.cell(row=r, column=1, value="Episode 1")
            ws.cell(row=r, column=2, value=name)
            ws.cell(row=r, column=4, value="00:00:02:00")
    wb.save(path)
    return path


@pytest.fixture
def old_version(tmp_path):
    return _write(tmp_path / "old.xlsx", [
        ("EP01 - Master XML", ["GETTYIMAGES-111", "AP-OLD-1", "MOVED-CLIP", "GONE-CLIP"]),
        ("AP", ["AP-OLD-1", "GONE-CLIP"]),
        ("Getty videos", ["GETTYIMAGES-111"]),
        ("Getty pics", ["GETTYIMAGES-777"]),
        ("3rd parties", ["MOVED-CLIP"]),
    ])


@pytest.fixture
def new_version(tmp_path):
    return _write(tmp_path / "new.xlsx", [
        ("Worksheet", ["GettyImages-111.mov", "AP-OLD-1", "MOVED-CLIP", "GettyImages-999", "AP-NEW-1", "GETTYIMAGES-777"]),
        ("AP", ["AP-OLD-1", "AP-NEW-1", "MOVED-CLIP"]),
        ("Getty Videos", ["GettyImages-111.mov", "GettyImages-999", "GettyImages-999"]),
        ("Getty Stills", ["GETTYIMAGES-777"]),
        ("3rd parties", []),
    ])


@pytest.fixture
def result(old_version, new_version, tmp_path):
    out = tmp_path / "report.xlsx"
    report = compare_workbooks(str(old_version), str(new_version), str(out))
    return load_workbook(out), report


def _rows(ws):
    return [(ws.cell(row=r, column=1).value, ws.cell(row=r, column=3).value)
            for r in range(2, ws.max_row + 1)]


def test_summary_is_first_and_lists_both_sides_counts(result):
    wb, _ = result
    ws = wb["Summary"]
    assert wb.sheetnames[0] == "Summary"
    assert [ws.cell(row=1, column=c).value for c in range(1, 8)] == [
        "Category", "Clips in old", "Clips in new", "Added", "Removed",
        "Moved in (from another old sheet)", "Moved out (to another new sheet)"]
    assert [ws.cell(row=2, column=c).value for c in range(1, 8)] == ["AP", 2, 3, 1, 1, 1, 0]


def test_summary_has_a_total_row(result):
    wb, report = result
    ws = wb["Summary"]
    total_row = len(report) + 3
    assert ws.cell(row=total_row, column=1).value == "TOTAL"
    assert ws.cell(row=total_row, column=4).value == sum(len(c["added"]) for c in report.values())


def test_clip_on_no_old_sheet_is_added(result):
    _, report = result
    assert report["Getty Videos"]["added"] == ["GettyImages-999"]
    assert report["AP"]["added"] == ["AP-NEW-1"]


def test_clip_that_only_changed_category_is_moved_not_added(result):
    _, report = result
    assert report["AP"]["moved_in"] == ["MOVED-CLIP"]
    assert report["AP"]["added"] == ["AP-NEW-1"]
    assert report["3rd parties"]["moved_out"] == ["MOVED-CLIP"]


def test_clip_on_no_new_sheet_is_removed(result):
    _, report = result
    assert report["AP"]["removed"] == ["GONE-CLIP"]


def test_getty_pics_pairs_with_getty_stills(result):
    _, report = result
    # Same clip, different sheet name across versions: not a diff at all.
    assert report["Getty Stills"] == {
        "old_clips": 1, "new_clips": 1, "old_rows": 1, "new_rows": 1,
        "added": [], "removed": [], "moved_in": [], "moved_out": []}


def test_case_and_extension_differences_are_not_diffs(result):
    _, report = result
    # "GETTYIMAGES-111" (old) vs "GettyImages-111.mov" (new).
    assert report["Getty Videos"]["added"] == ["GettyImages-999"]
    assert report["Getty Videos"]["removed"] == []


def test_backup_tabs_get_no_category_row(result):
    wb, report = result
    assert "Worksheet" not in report and "EP01 - Master XML" not in report
    assert not any(title.startswith("Worksheet") for title in wb.sheetnames)


def test_added_sheet_carries_rows_with_a_status_column(result):
    wb, _ = result
    ws = wb["AP added"]
    assert ws.cell(row=1, column=1).value == "Status"
    assert ws.cell(row=1, column=3).value == "Name"
    assert _rows(ws) == [("added", "AP-NEW-1"), ("moved from 3rd parties", "MOVED-CLIP")]


def test_removed_sheet_comes_from_the_old_file(result):
    wb, _ = result
    assert _rows(wb["AP removed"]) == [("removed", "GONE-CLIP")]
    assert _rows(wb["3rd parties removed"]) == [("moved to AP", "MOVED-CLIP")]


def test_every_use_of_an_added_clip_is_listed_by_default(result):
    wb, _ = result
    assert _rows(wb["Getty Videos added"]) == [("added", "GettyImages-999"), ("added", "GettyImages-999")]


def test_unique_lists_an_added_clip_once(old_version, new_version, tmp_path):
    out = tmp_path / "unique.xlsx"
    compare_workbooks(str(old_version), str(new_version), str(out), unique=True)
    assert _rows(load_workbook(out)["Getty Videos added"]) == [("added", "GettyImages-999")]


def test_categories_with_no_diff_get_no_sheet(result):
    wb, _ = result
    assert wb.sheetnames == [
        "Summary", "AP added", "AP removed", "Getty Videos added", "3rd parties removed"]


def test_case_sensitive_treats_recased_names_as_diffs(old_version, new_version, tmp_path):
    out = tmp_path / "cs.xlsx"
    report = compare_workbooks(str(old_version), str(new_version), str(out), case_sensitive=True)
    assert "GettyImages-111.mov" in report["Getty Videos"]["added"]
    assert "GETTYIMAGES-111" in report["Getty Videos"]["removed"]


def test_compare_output_styling_is_black_white(old_version, new_version, tmp_path):
    out = tmp_path / "styled.xlsx"
    compare_workbooks(str(old_version), str(new_version), str(out))
    wb = load_workbook(out)

    for t in wb.sheetnames:
        ws = wb[t]
        h = ws.cell(row=1, column=1)
        assert h.fill.fgColor.rgb == "FF000000", t
        assert h.font.bold and h.font.color.rgb == "FFFFFFFF", t

    row_sheets = [t for t in wb.sheetnames
                  if t.endswith("added") or t.endswith("removed")]
    assert row_sheets
    assert 10 <= wb[row_sheets[0]].column_dimensions["A"].width <= 60

    # Summary keeps its fixed first-column width (not passed to style_output_sheets).
    assert wb["Summary"].column_dimensions["A"].width == 24


def test_status_column_stays_bold_after_styling(old_version, new_version, tmp_path):
    # The Status column (A) is deliberately bold; style_output_sheets' uniform
    # data font runs over it, so compare_workbooks must re-bold it afterwards.
    out = tmp_path / "styled.xlsx"
    compare_workbooks(str(old_version), str(new_version), str(out))
    wb = load_workbook(out)
    ws = wb["AP added"]
    assert ws.cell(row=2, column=1).value  # a real status row
    assert ws.cell(row=2, column=1).font.bold is True


def test_grouped_totals_block_is_not_counted_as_a_clip(old_version, tmp_path):
    new = _write(tmp_path / "grouped.xlsx", [("AP", ["AP-OLD-1", "GONE-CLIP"])])
    wb = load_workbook(new)
    wb["AP"].cell(row=4, column=2, value="Total clips")
    wb["AP"].cell(row=5, column=2, value="Total Seconds")
    wb.save(new)
    report = compare_workbooks(str(old_version), str(new), str(tmp_path / "out.xlsx"))
    assert report["AP"] == {
        "old_clips": 2, "new_clips": 2, "old_rows": 2, "new_rows": 2,
        "added": [], "removed": [], "moved_in": [], "moved_out": []}
