import io

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from web.auth import hash_password
from web.main import app


def _client(monkeypatch):
    monkeypatch.setenv("BIGMEDIA_WEB_PASSWORD_HASH", hash_password("pw"))
    c = TestClient(app)
    c.post("/login", data={"password": "pw"})
    return c


def _master():
    wb = Workbook()
    ws = wb.active
    ws.title = "Master"
    ws.append(["Media File", "Clip Duration"])
    for name in [
        "BM1234_liberation.mxf",                 # AP
        "BM5678_parade.mxf",                     # AP
        "045_SO_EP18_01_3DExplainer_TXLS.mov",   # GFX
        "shutterstock_777.mp4",                  # Shutterstock
        "random_unclassified_xyz.mov",           # 3rd parties
    ]:
        ws.append([name, "00:00:10:00"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_analyze_then_run_sort(monkeypatch):
    client = _client(monkeypatch)
    content = _master()

    # 1. Analyze with the real (non-default) name column.
    analyze = client.post(
        "/commands/sort/analyze",
        data={"name_column": "Media File"},
        files={"files": ("master.xlsx", content, "application/octet-stream")},
    ).json()
    assert analyze["headers"] == ["Media File", "Clip Duration"]
    assert analyze["annotations"]["categories"]["AP"] == 2
    assert set(analyze["suggestions"]["categories"]) >= {"AP", "GFX", "Shutterstock"}
    assert not analyze["warnings"]

    # 2. User keeps only AP + GFX, then runs.
    run = client.post(
        "/commands/sort",
        data={"name_column": "Media File", "categories": ["AP", "GFX"]},
        files={"files": ("master.xlsx", content, "application/octet-stream")},
    )
    assert run.status_code == 200
    wb = load_workbook(io.BytesIO(run.content))
    assert "AP" in wb.sheetnames and "GFX" in wb.sheetnames
    assert "Shutterstock" not in wb.sheetnames
    ap = wb["AP"]
    assert ap.max_row == 3  # header + 2 AP clips
