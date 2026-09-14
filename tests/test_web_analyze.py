from openpyxl import Workbook

from web.analyze import run_analysis
from web.commands import COMMANDS, CommandSpec, FieldSpec


def _sheet(path, headers, names):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for n in names:
        ws.append([n])
    wb.save(path)
    return str(path)


def _multi_sheet(path, sheets):
    """sheets: dict of {title: [header_row, *data_rows]}, first key becomes
    the active sheet (matches openpyxl's default: first-created = active)."""
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title=title)
        for row in rows:
            ws.append(row)
    wb.save(path)
    return str(path)


def test_generic_only_command_returns_headers_and_sheets(tmp_path):
    path = _sheet(tmp_path / "m.xlsx", ["Clip Name"], ["x.mov"])
    result = run_analysis(COMMANDS["dedupe"], [path], {})
    assert result["headers"] == ["Clip Name"]
    assert result["sheets"][0]["name"]
    # Generic column auto-detect (Fix 3) fills name_column; no per-command
    # analyzer runs, so nothing else is suggested and no annotations.
    assert result["suggestions"] == {"name_column": "Clip Name"}
    assert result["annotations"] == {}


def test_sort_merges_analyzer_suggestions_and_annotations(tmp_path):
    path = _sheet(tmp_path / "m.xlsx", ["Clip Name"], ["BM1234_x.mxf", "random_xyz.mov"])
    result = run_analysis(COMMANDS["sort"], [path], {"name_column": "Clip Name"})
    assert "AP" in result["suggestions"]["categories"]
    assert result["annotations"]["categories"]["AP"] == 1


def test_sort_analyzer_uses_form_name_column(tmp_path):
    path = _sheet(tmp_path / "m.xlsx", ["Media File"], ["BM1234_x.mxf"])
    result = run_analysis(COMMANDS["sort"], [path], {"name_column": "Media File"})
    assert result["annotations"]["categories"]["AP"] == 1
    assert not result["warnings"]


def test_sort_analyzer_warning_is_surfaced(tmp_path):
    path = _sheet(tmp_path / "m.xlsx", ["Media File"], ["BM1234_x.mxf"])
    result = run_analysis(COMMANDS["sort"], [path], {"name_column": "Clip Name"})
    assert result["warnings"]


# --- Fix 3: name/duration column auto-detect (spec decision 4) --------------

def test_autodetect_name_column_clip_name(tmp_path):
    path = _sheet(tmp_path / "m.xlsx", ["Clip Name", "Clip Duration"], ["x.mov"])
    result = run_analysis(COMMANDS["dedupe"], [path], {})
    assert result["suggestions"]["name_column"] == "Clip Name"


def test_autodetect_name_column_falls_back_to_name(tmp_path):
    path = _sheet(tmp_path / "m.xlsx", ["Name"], ["x.mov"])
    result = run_analysis(COMMANDS["dedupe"], [path], {})
    assert result["suggestions"]["name_column"] == "Name"


def test_autodetect_duration_column(tmp_path):
    path = _sheet(tmp_path / "m.xlsx", ["Clip Name", "Duration"], ["x.mov"])
    result = run_analysis(COMMANDS["dedupe"], [path], {})
    assert result["suggestions"]["duration_column"] == "Duration"


def test_autodetect_absent_when_no_nameish_column(tmp_path):
    path = _sheet(tmp_path / "m.xlsx", ["Media File", "Length"], ["x.mov"])
    result = run_analysis(COMMANDS["dedupe"], [path], {})
    assert "name_column" not in result["suggestions"]
    assert "duration_column" not in result["suggestions"]


def test_per_command_analyzer_still_overrides_autodetect(tmp_path):
    # Sort's analyzer emits suggestions.categories only, not name_column,
    # so the auto-detected name_column must survive alongside it.
    path = _sheet(tmp_path / "m.xlsx", ["Clip Name"], ["BM1234_x.mxf"])
    result = run_analysis(COMMANDS["sort"], [path], {"name_column": "Clip Name"})
    assert result["suggestions"]["name_column"] == "Clip Name"
    assert "categories" in result["suggestions"]


# --- Fix 6: _analyzer_kwargs type-safety + param guards --------------------

def test_analyzer_param_not_in_fields_is_not_passed(tmp_path, monkeypatch):
    path = _sheet(tmp_path / "m.xlsx", ["Clip Name"], ["BM1234_x.mxf"])
    seen = {}

    def fake_analyze(paths, name_column="Clip Name", files=None, bogus=None):
        seen["files"] = files
        seen["bogus"] = bogus
        seen["name_column"] = name_column
        return {"suggestions": {}}

    spec = CommandSpec(
        slug="fake", title="Fake", upload_mode="batch", func=lambda *a, **k: None,
        output_suffix="x", analyze=fake_analyze,
        fields=[FieldSpec("name_column", "Filename column", "text", "Clip Name")],
    )
    run_analysis(spec, [path], {"name_column": "Clip Name", "files": "SHOULD_NOT_LEAK", "bogus": "x"})
    assert seen["files"] is None
    assert seen["bogus"] is None
    assert seen["name_column"] == "Clip Name"


# --- sheet_source: sheet-scoped headers/warnings/suggestions ---------------

def test_getty_ids_headers_scoped_to_sheet_name_not_active_sheet(tmp_path):
    path = _multi_sheet(tmp_path / "ep1.xlsx", {
        "COST": [["EP1"]],
        "Getty Videos": [["Clip Name", "Seconds"], ["x.mov", 6]],
    })
    result = run_analysis(COMMANDS["getty-ids"], [path], {"sheet_name": "Getty Videos", "stills_sheet_name": "Getty Stills"})
    assert result["headers"] == ["Clip Name", "Seconds"]
    assert result["suggestions"]["name_column"] == "Clip Name"


def test_getty_ids_stills_sheet_missing_everywhere_suggests_none_no_warning(tmp_path):
    path = _multi_sheet(tmp_path / "ep1.xlsx", {"Getty Videos": [["Clip Name"]]})
    result = run_analysis(COMMANDS["getty-ids"], [path], {"sheet_name": "Getty Videos", "stills_sheet_name": "Getty Stills"})
    assert result["suggestions"]["stills_sheet_name"] == ""
    assert not result["warnings"]


def test_getty_ids_missing_required_video_sheet_warns(tmp_path):
    path = _multi_sheet(tmp_path / "ep1.xlsx", {"COST": [["EP1"]]})
    result = run_analysis(COMMANDS["getty-ids"], [path], {"sheet_name": "Getty Videos", "stills_sheet_name": "Getty Stills"})
    assert any("Getty Videos" in w and "not found" in w for w in result["warnings"])


def test_getty_ids_sheet_name_suggests_case_insensitive_match(tmp_path):
    path = _multi_sheet(tmp_path / "ep1.xlsx", {"getty videos": [["Clip Name"]]})
    result = run_analysis(COMMANDS["getty-ids"], [path], {})
    assert result["suggestions"]["sheet_name"] == "getty videos"


def test_getty_ids_stills_sheet_suggests_via_alias(tmp_path):
    path = _multi_sheet(tmp_path / "ep1.xlsx", {"Getty Videos": [["Clip Name"]], "Getty pics": [["Clip Name"]]})
    result = run_analysis(COMMANDS["getty-ids"], [path], {})
    assert result["suggestions"]["stills_sheet_name"] == "Getty pics"


def test_fu_grid_headers_scoped_to_selected_sheet(tmp_path):
    path = _multi_sheet(tmp_path / "ep1.xlsx", {
        "3rd parties": [["Clip Name", "Clip Duration"], ["x.mov", 5]],
        "COST": [["EP1"]],
    })
    result = run_analysis(COMMANDS["fu-grid"], [path], {"sheet": "3rd parties"})
    assert result["headers"] == ["Clip Name", "Clip Duration"]


def test_fu_grid_sheet_field_suggests_default_when_present(tmp_path):
    path = _multi_sheet(tmp_path / "ep1.xlsx", {"3rd parties": [["Clip Name"]]})
    result = run_analysis(COMMANDS["fu-grid"], [path], {})
    assert result["suggestions"]["sheet"] == "3rd parties"


def test_group_headers_scoped_to_checked_sheets_union(tmp_path):
    path = _multi_sheet(tmp_path / "ep1.xlsx", {
        "AP": [["Clip Name"]],
        "Getty Videos": [["Clip Name", "Clip Duration"]],
        "COST": [["EP1"]],
    })
    result = run_analysis(COMMANDS["group"], [path], {"sheets": ["AP", "Getty Videos"]})
    assert set(result["headers"]) == {"Clip Name", "Clip Duration"}


def test_group_sheets_checklist_suggests_default_sheets_present(tmp_path):
    path = _multi_sheet(tmp_path / "ep1.xlsx", {
        "AP": [["Clip Name"]],
        "Getty Videos": [["Clip Name"]],
        "COST": [["EP1"]],
    })
    result = run_analysis(COMMANDS["group"], [path], {})
    assert set(result["suggestions"]["sheets"]) == {"AP", "Getty Videos"}
    assert "COST" not in result["suggestions"]["sheets"]


def test_group_sheets_checklist_matches_stills_alias(tmp_path):
    path = _multi_sheet(tmp_path / "ep1.xlsx", {"Getty pics": [["Clip Name"]]})
    result = run_analysis(COMMANDS["group"], [path], {})
    assert result["suggestions"]["sheets"] == ["Getty pics"]


def test_sort_dedupe_headers_still_active_sheet_based(tmp_path):
    # No sheet_source anywhere for these two -> untouched code path.
    path = _multi_sheet(tmp_path / "ep1.xlsx", {
        "COST": [["EP1"]],
        "Getty Videos": [["Clip Name"]],
    })
    result = run_analysis(COMMANDS["dedupe"], [path], {})
    assert result["headers"] == ["EP1"]
