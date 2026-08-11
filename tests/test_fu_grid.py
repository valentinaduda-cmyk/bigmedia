"""Tests for the FU grid command.

The sample_third_parties fixture (tests/conftest.py) stands in for a
sorted workbook's "3rd parties" sheet, with rows deliberately out of
sequence-in order so the "earliest TC IN wins" rule is exercised.
"""
import pytest
from openpyxl import load_workbook

from bigmedia.fu_grid import fu_grid_workbook


@pytest.fixture
def grid(sample_third_parties, tmp_path):
    """The "FU grid" sheet produced from the fixture, plus the counts dict."""
    out = tmp_path / "fu_grid.xlsx"
    counts = fu_grid_workbook(str(sample_third_parties), str(out))
    return load_workbook(out), counts


def _row_values(ws, row):
    return [ws.cell(row=row, column=c).value for c in range(1, 14)]


def test_output_has_backup_sheet_first_then_grid(grid):
    wb, _ = grid
    assert wb.sheetnames == ["3rd parties", "FU grid"]


def test_backup_sheet_is_a_verbatim_copy(grid, sample_third_parties):
    wb, _ = grid
    ws_src = load_workbook(sample_third_parties)["3rd parties"]
    ws_out = wb["3rd parties"]
    assert (ws_out.max_row, ws_out.max_column) == (ws_src.max_row, ws_src.max_column)
    for r in range(1, ws_src.max_row + 1):
        for c in range(1, ws_src.max_column + 1):
            assert ws_out.cell(row=r, column=c).value == ws_src.cell(row=r, column=c).value


def test_header_row(grid):
    wb, _ = grid
    assert _row_values(wb["FU grid"], 1) == [
        "TC IN", "TC OUT", "CLIP DURATION", "TOTAL USES",
        "TOTAL DURATION if multiple uses", "SOURCE DURATION", "SCREENSHOTS",
        "URL LINK", "PREVIOUS LEGAL CHECK", "SOURCE", "CLIP NAME",
        "FINAL FU LEGAL NOTE", "PREVIEW: ",
    ]


def test_groups_by_exact_clip_name_not_dedup_key(grid):
    wb, _ = grid
    ws = wb["FU grid"]
    names = [ws.cell(row=r, column=11).value for r in range(2, ws.max_row + 1)]
    # ".mp4" and ".mov" are the same clip to dedup_key, but not here.
    assert names == ["Evan Fairbanks WTC.mp4", "Evan Fairbanks WTC.mov", "Firehouse raw.mp4"]


def test_rows_ordered_by_earliest_tc_in(grid):
    wb, _ = grid
    ws = wb["FU grid"]
    tc_ins = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
    assert tc_ins == ["01:00:05:00", "01:00:20:00", "01:00:30:00"]


def test_surviving_row_is_the_earliest_use(grid):
    wb, _ = grid
    # The 3-use group's earliest use is 01:00:05:00 -- its TC OUT, clip
    # duration and source duration must all come from that same row.
    assert _row_values(wb["FU grid"], 2)[:6] == [
        "01:00:05:00", "01:00:05:10", "00:00:00:10", 3, "00:00:01:20", "00:00:00:21",
    ]


def test_total_duration_sums_the_group_with_frame_carry(grid):
    wb, _ = grid
    # 00:00:00:10 + 00:00:00:20 + 00:00:00:15 = 45 frames = 1s 20f at 25 fps.
    assert wb["FU grid"].cell(row=2, column=5).value == "00:00:01:20"


def test_total_duration_repeats_the_clip_duration_for_single_use_clips(grid):
    wb, _ = grid
    ws = wb["FU grid"]
    assert ws.cell(row=3, column=4).value == 1
    assert ws.cell(row=3, column=5).value == ws.cell(row=3, column=3).value == "00:00:02:00"


def test_manual_columns_are_left_empty(grid):
    wb, _ = grid
    ws = wb["FU grid"]
    for row in range(2, ws.max_row + 1):
        for col in (7, 8, 9, 10, 12, 13):
            assert ws.cell(row=row, column=col).value is None


def test_counts(grid):
    _, counts = grid
    assert counts == {"rows_in": 5, "unique_clips": 3, "multi_use_clips": 1}


def test_data_row_fills(grid):
    wb, _ = grid
    ws = wb["FU grid"]
    assert ws.cell(row=2, column=1).fill.fgColor.rgb == "FFF4CCCC"
    assert ws.cell(row=2, column=9).fill.fgColor.rgb == "FFD9EAD3"


def test_header_styling(grid):
    wb, _ = grid
    cell = wb["FU grid"].cell(row=1, column=1)
    assert cell.fill.fgColor.rgb == "FFCFE2F3"
    assert cell.font.bold
    assert cell.alignment.horizontal == "center"
    assert cell.alignment.wrap_text


def test_totals_columns_are_bold(grid):
    wb, _ = grid
    ws = wb["FU grid"]
    assert [ws.cell(row=2, column=c).font.bold for c in (3, 4, 5, 6)] == [False, True, True, True]


def test_blank_clip_names_are_skipped(sample_third_parties, tmp_path):
    wb = load_workbook(sample_third_parties)
    ws = wb["3rd parties"]
    ws.cell(row=ws.max_row + 1, column=1, value="V3")
    ws.cell(row=ws.max_row, column=2, value="   ")
    wb.save(sample_third_parties)

    out = tmp_path / "fu_grid.xlsx"
    counts = fu_grid_workbook(str(sample_third_parties), str(out))
    assert counts == {"rows_in": 5, "unique_clips": 3, "multi_use_clips": 1}


def test_blank_duration_counts_as_zero(sample_third_parties, tmp_path):
    wb = load_workbook(sample_third_parties)
    ws = wb["3rd parties"]
    ws.cell(row=2, column=6).value = None  # one of the 3-use group's rows
    wb.save(sample_third_parties)

    out = tmp_path / "fu_grid.xlsx"
    fu_grid_workbook(str(sample_third_parties), str(out))
    # 00:00:00:10 + (blank) + 00:00:00:15 = 25 frames = exactly 1s.
    assert load_workbook(out)["FU grid"].cell(row=2, column=5).value == "00:00:01:00"


def test_sheet_matched_case_insensitively(sample_third_parties, tmp_path):
    wb = load_workbook(sample_third_parties)
    # Two steps: openpyxl uniquifies a rename that collides with the
    # sheet's own current title.
    wb["3rd parties"].title = "tmp"
    wb["tmp"].title = "3RD PARTIES"
    wb.save(sample_third_parties)

    out = tmp_path / "fu_grid.xlsx"
    fu_grid_workbook(str(sample_third_parties), str(out))
    assert load_workbook(out).sheetnames == ["3RD PARTIES", "FU grid"]


def test_missing_sheet_raises_naming_available_sheets(sample_third_parties, tmp_path):
    out = tmp_path / "fu_grid.xlsx"
    with pytest.raises(ValueError) as excinfo:
        fu_grid_workbook(str(sample_third_parties), str(out), sheet="Nope")
    assert "Nope" in str(excinfo.value)
    assert "3rd parties" in str(excinfo.value)


def test_empty_source_sheet_still_writes_a_header(sample_third_parties, tmp_path):
    wb = load_workbook(sample_third_parties)
    ws = wb["3rd parties"]
    ws.delete_rows(2, ws.max_row)
    wb.save(sample_third_parties)

    out = tmp_path / "fu_grid.xlsx"
    counts = fu_grid_workbook(str(sample_third_parties), str(out))
    assert counts == {"rows_in": 0, "unique_clips": 0, "multi_use_clips": 0}
    assert load_workbook(out)["FU grid"].max_row == 1


def test_fps_affects_the_total_duration_sum(sample_third_parties, tmp_path):
    out = tmp_path / "fu_grid.xlsx"
    fu_grid_workbook(str(sample_third_parties), str(out), fps=50)
    # 10 + 20 + 15 = 45 frames, under one second at 50 fps.
    assert load_workbook(out)["FU grid"].cell(row=2, column=5).value == "00:00:00:45"
