from openpyxl import Workbook, load_workbook
from bigmedia.dedupe import dedup_key, dedupe_workbook
from bigmedia.xlsx_utils import find_column


def test_dedup_key_strips_extension_and_copy_marker():
    assert dedup_key("Clip_A.mov") == "Clip_A"
    assert dedup_key("Clip_A (1).mxf") == "Clip_A"
    assert dedup_key("Clip_A") == "Clip_A"


def test_dedup_key_strips_trailing_number_tag_after_extension():
    # A real duplicate pair: same clip, one with a stray " 25" tag appended
    # after the extension (e.g. a frame-rate/version annotation).
    assert dedup_key("GettyImages-627-122.mov 25") == dedup_key("GettyImages-627-122.mov")
    assert dedup_key("GettyImages-627-122.mov 25") == "GettyImages-627-122"
    # Without an extension, a trailing " <digits>" is NOT noise -- it can be
    # a meaningful part of the id (e.g. "GettyImages-805-74 2"), so it must
    # be left alone.
    assert dedup_key("GettyImages-805-74 2") == "GettyImages-805-74 2"


def test_dedup_key_strips_known_codec_suffix():
    # Same clip, one re-exported with an Apple ProRes codec tag baked into
    # the filename before the extension -- not a different clip.
    assert dedup_key("GettyImages-1479649397_Apple_ProRes_422.mov") == dedup_key(
        "GettyImages-1479649397.mov"
    )
    assert dedup_key("GettyImages-1479649397_Apple_ProRes_422.mov") == "GettyImages-1479649397"
    # A real id that merely contains digits after an underscore, with no
    # codec keyword, must NOT be stripped.
    assert dedup_key("GettyImages-356-30_1234.mov") == "GettyImages-356-30_1234"


def test_dedup_key_strips_getty_upscale_suffix():
    # Same Getty clip, re-exported/upscaled with an "_S000_upscaleNN" or bare
    # "_upscaleNN" tag baked into the filename before the extension.
    assert dedup_key("GettyImages-650878972_S000_upscale01.mov") == "GettyImages-650878972"
    assert dedup_key("GettyImages-650878972_S001_upscale02.mov") == "GettyImages-650878972"
    assert dedup_key("GettyImages-650878972_S002_upscale03.mov") == "GettyImages-650878972"
    assert dedup_key("GettyImages-650878972_upscale01.mov") == "GettyImages-650878972"


def test_dedup_key_strips_known_post_process_suffix():
    # Same Getty clip, re-exported through a post-process tool (Avid retime,
    # denoise, deflicker, NTSC standards conversion) with the tool's tag
    # baked into the filename before the extension -- not a different clip.
    assert dedup_key("GettyImages-551409951_AvidRetime-10383301") == "GettyImages-551409951"
    assert dedup_key("GettyImages-643936503_AvidRetime-10474212") == "GettyImages-643936503"
    assert dedup_key("GettyImages-502266528_Denoise_chr2") == "GettyImages-502266528"
    assert dedup_key("GettyImages-1335217652.mp4 25") == dedup_key("GettyImages-1335217652_Denoise")
    assert dedup_key("GettyImages-1335217652_Denoise") == "GettyImages-1335217652"
    assert dedup_key("GettyImages-2247246190_Denoise") == "GettyImages-2247246190"
    assert dedup_key("GettyImages-673768296_Deflicker") == "GettyImages-673768296"
    assert dedup_key("GettyImages-85263020_NTSC") == "GettyImages-85263020"


def test_dedup_key_strips_getty_render_prob4_comp_junk():
    # Getty QC/review-export pipelines bake "Render"/"Render <N>" take
    # markers, "_prob4" problem-pass tags, and "_comp"/"_comp_<N>" composite
    # tags onto the same source id, re-appending the extension on every
    # pass -- not different clips.
    pairs = [
        ("GettyImages-1142748666.mov Render 1_prob4.mov", "GettyImages-1142748666.mov Render_prob4.mov"),
        ("GettyImages-1160882554.mov Render 1_prob4.mov Render.mov", "GettyImages-1160882554.mov Render_prob4.mov"),
        ("GettyImages-1171009582.mov Render.mov Render_1_prob4.mov", "GettyImages-1171009582_comp"),
        ("GettyImages-1200009115.mov", "GettyImages-1200009115.mov Render 1_prob4.mov Render.mov"),
        ("GettyImages-1200009155.mov Render 1_prob4.mov Render.mov", "GettyImages-1200009155.mov Render 2_prob4.mov Render.mov"),
        ("GettyImages-1200009155.mov Render 1_prob4.mov Render.mov", "GettyImages-1200009155.mov Render_prob4.mov Render.mov"),
        ("GettyImages-1262596964.mov Render 1_prob4.mov Render.mov", "GettyImages-1262596964.mov Render_prob4.mov Render.mov"),
        ("GettyImages-1316659261 Render.mov", "GettyImages-1316659261 Render 1.mov"),
        ("GettyImages-1442787190.mov Render 1_prob4.mov", "GettyImages-1442787190.mov Render_prob4.mov"),
        ("GettyImages-2064403302_comp", "GettyImages-2064403302_comp_2"),
        ("GettyImages-2190433054.mov Render 1_prob4.mov Render.mov", "GettyImages-2190433054.mov Render.mov Render_1_prob4.mov Render.mov"),
        ("GettyImages-3B2FBBF3_0377.mov", "GettyImages-3B2FBBF3_0377_comp"),
    ]
    for a, b in pairs:
        assert dedup_key(a) == dedup_key(b), f"{a!r} vs {b!r}"
    # Sanity: none of these collapsed to an empty/wrong id.
    assert dedup_key("GettyImages-1142748666.mov Render 1_prob4.mov") == "GettyImages-1142748666"
    assert dedup_key("GettyImages-2064403302_comp") == "GettyImages-2064403302"
    assert dedup_key("GettyImages-3B2FBBF3_0377_comp") == "GettyImages-3B2FBBF3_0377"
    # A trailing " <digits>" with no "Render"/extension is still meaningful,
    # not junk -- must not be swallowed by the new stripping.
    assert dedup_key("GettyImages-805-74 2") == "GettyImages-805-74 2"


def test_dedup_key_strips_new_post_process_tags():
    # New export/post-process tags baked onto the same source clip on
    # re-render -- not different clips. Each is only a trailing "_"-segment
    # (or the "-640_ADPP" resolution/tool block), stripped before/around the
    # extension and any " (N)" copy marker.
    assert dedup_key("GettyImages-1234567890_SM.mov") == "GettyImages-1234567890"
    assert dedup_key("GettyImages-1234567890_SM01.mov") == "GettyImages-1234567890"
    assert dedup_key("GettyImages-1234567890_SM02.MP4") == "GettyImages-1234567890"
    assert dedup_key("GettyImages-1234567890_AIUPSCALE.mov") == "GettyImages-1234567890"
    assert dedup_key("GettyImages-1234567890_APPLEPRORESHQ") == "GettyImages-1234567890"
    assert dedup_key("GettyImages-1234567890_APPLEPRORESHQ.MOV") == "GettyImages-1234567890"
    assert dedup_key("GettyImages-1234567890_DFR") == "GettyImages-1234567890"
    assert dedup_key("GettyImages-1234567890_DFR01") == "GettyImages-1234567890"
    assert dedup_key("GettyImages-1234567890_DFR_OK") == "GettyImages-1234567890"
    assert dedup_key("GettyImages-1234567890_DFR01_OK") == "GettyImages-1234567890"
    assert dedup_key("GettyImages-1234567890_DFR02_OK.mov") == "GettyImages-1234567890"


def test_dedup_key_strips_adpp_block_and_bracketed_number():
    # "-640_ADPP" is a fixed block; the number in brackets before the
    # extension is a copy marker and varies (with or without a space).
    assert dedup_key("GettyImages-2217807802-640_ADPP.MP4") == "GettyImages-2217807802"
    assert dedup_key("GettyImages-2217807802-640_ADPP (3).MP4") == "GettyImages-2217807802"
    assert dedup_key("GettyImages-2217807802-640_ADPP(4).MP4") == "GettyImages-2217807802"
    assert dedup_key("GettyImages-2217807802-640_ADPP (3).MP4") == dedup_key(
        "GettyImages-2217807802.mov"
    )


def test_dedup_key_strips_combinations_of_tags_in_any_order():
    base = "GettyImages-1234567890"
    for name in [
        "GettyImages-1234567890_SM01_APPLEPRORESHQ.MOV",
        "GettyImages-1234567890_APPLEPRORESHQ_AIUPSCALE.mov",
        "GettyImages-1234567890_SM_DFR_OK.mov",
        "GettyImages-1234567890_DFR01_OK_SM02.MP4",
        "GettyImages-1234567890-640_ADPP_AIUPSCALE (2).MP4",
        "GettyImages-1234567890_Apple_ProRes_422_SM01.mov",
    ]:
        assert dedup_key(name) == base, name


def test_dedup_key_new_tags_do_not_overreach():
    # A real id segment that merely looks tag-like must survive: "_SM" /
    # "_DFR" only strip as a trailing segment, and "-640_ADPP" is literal.
    assert dedup_key("GettyImages-356-30_1234.mov") == "GettyImages-356-30_1234"
    assert dedup_key("GettyImages-805-74 2") == "GettyImages-805-74 2"
    assert dedup_key("GettyImages-1B010728_t010.mov") == "GettyImages-1B010728_t010"


def test_dedupe_workbook_splits_kept_and_dropped(sample_master, tmp_path):
    out_path = tmp_path / "out.xlsx"
    result = dedupe_workbook(str(sample_master), str(out_path))
    # fixture has 7 rows, one of which is a "(1)" duplicate of another
    assert result["kept"] == 6
    assert result["dropped"] == 1


def test_dedupe_workbook_fills_source_reel_name(sample_master, tmp_path):
    out_path = tmp_path / "out.xlsx"
    dedupe_workbook(str(sample_master), str(out_path))

    wb = load_workbook(out_path)
    for title in ["Deduped", "Duplicates"]:
        ws = wb[title]
        source_col_idx = find_column(ws, "Source Reel Name")
        for r in range(2, ws.max_row + 1):
            assert ws.cell(row=r, column=source_col_idx).value == title

    # Backup rows are labeled with whichever sheet they ended up on.
    src = load_workbook(sample_master).active
    backup = wb["Worksheet"]
    source_col_idx = find_column(src, "Source Reel Name")
    duplicate_row = 8  # the "(1)" copy of row 2, per the fixture
    assert backup.cell(row=duplicate_row, column=source_col_idx).value == "Duplicates"
    assert backup.cell(row=2, column=source_col_idx).value == "Deduped"


def test_dedupe_workbook_falls_back_to_name_column(tmp_path):
    # Some real delivery files title the clip-name column "Name" instead of
    # "Clip Name" (the default). dedupe_workbook should still find it
    # without requiring name_column to be passed explicitly.
    src_path = tmp_path / "src.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Name", "Source Reel Name"])
    ws.append(["Clip_A.mov", ""])
    ws.append(["Clip_A (1).mxf", ""])
    wb.save(src_path)

    out_path = tmp_path / "out.xlsx"
    result = dedupe_workbook(str(src_path), str(out_path))
    assert result["kept"] == 1
    assert result["dropped"] == 1


def test_dedupe_output_sheets_are_styled(sample_master, tmp_path):
    from openpyxl import load_workbook
    out = tmp_path / "out.xlsx"
    dedupe_workbook(str(sample_master), str(out))
    wb = load_workbook(str(out))
    for t in ("Deduped", "Duplicates"):
        ws = wb[t]
        h = ws.cell(row=1, column=1)
        assert h.fill.fgColor.rgb == "FF000000"
        assert h.font.bold and h.font.color.rgb == "FFFFFFFF"
        assert 10 <= ws.column_dimensions["A"].width <= 60
    # backup untouched
    bak = wb["Worksheet"]
    src = load_workbook(str(sample_master)).active
    assert bak.cell(row=1, column=1).value == src.cell(row=1, column=1).value
