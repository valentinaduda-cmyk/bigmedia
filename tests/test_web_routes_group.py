import io

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from web.auth import hash_password
from web.main import app


def _logged_in_client(monkeypatch):
    monkeypatch.setenv("BIGMEDIA_WEB_PASSWORD_HASH", hash_password("pw"))
    client = TestClient(app)
    client.post("/login", data={"password": "pw"})
    return client


def _sorted_workbook_bytes():
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Getty Videos")
    ws.append(["Clip Name", "Clip Duration"])
    ws.append(["GettyImages-12345.mov", "00:00:00:05"])
    ap = wb.create_sheet("AP")
    ap.append(["Clip Name", "Clip Duration"])
    ap.append(["BM1234_x.mxf", "00:00:00:05"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_group_sheets_checkboxes_group_only_checked_sheets(monkeypatch):
    client = _logged_in_client(monkeypatch)
    files = {"files": ("ep1.xlsx", _sorted_workbook_bytes(), "application/octet-stream")}
    response = client.post(
        "/commands/group",
        data={"sheets": ["Getty Videos"], "sheets_present": "1", "fps": "25"},
        files=files,
    )
    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.content))
    header_row = [c.value for c in next(wb["Getty Videos"].iter_rows(min_row=1, max_row=1))]
    assert "Total Duration" in header_row
    ap_header_row = [c.value for c in next(wb["AP"].iter_rows(min_row=1, max_row=1))]
    assert "Total Duration" not in ap_header_row  # AP wasn't checked -> untouched


def test_group_every_checkbox_unchecked_with_marker_groups_nothing(monkeypatch):
    client = _logged_in_client(monkeypatch)
    files = {"files": ("ep1.xlsx", _sorted_workbook_bytes(), "application/octet-stream")}
    response = client.post(
        "/commands/group",
        data={"sheets_present": "1", "fps": "25"},
        files=files,
    )
    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.content))
    header_row = [c.value for c in next(wb["Getty Videos"].iter_rows(min_row=1, max_row=1))]
    assert "Total Duration" not in header_row


def test_group_no_marker_plain_form_keeps_default_sheets(monkeypatch):
    client = _logged_in_client(monkeypatch)
    files = {"files": ("ep1.xlsx", _sorted_workbook_bytes(), "application/octet-stream")}
    response = client.post("/commands/group", data={"fps": "25"}, files=files)
    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.content))
    header_row = [c.value for c in next(wb["Getty Videos"].iter_rows(min_row=1, max_row=1))]
    assert "Total Duration" in header_row  # DEFAULT_SHEETS still applies
