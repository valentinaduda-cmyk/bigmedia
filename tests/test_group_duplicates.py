from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from bigmedia.group_duplicates import group_duplicates_workbook
from bigmedia.xlsx_utils import find_column


def test_group_duplicates_reorders_and_sums_getty_videos(sample_sorted_master, tmp_path):
    out_path = tmp_path / "out.xlsx"
    results = group_duplicates_workbook(str(sample_sorted_master), str(out_path))

    wb = load_workbook(out_path)
    ws = wb["Getty Videos"]

    name_col = find_column(ws, "Clip Name")
    dur_col = find_column(ws, "Clip Duration")
    total_col = dur_col + 1  # inserted right after Clip Duration, not at the end
    seconds_col = dur_col + 2

    assert ws.cell(row=1, column=total_col).value == "Total Duration"
    assert ws.cell(row=1, column=seconds_col).value == "Seconds"
    # Columns after Clip Duration got pushed right by two.
    assert ws.cell(row=1, column=seconds_col + 1).value == "Source Reel Name"

    # Rows are reordered into ascending-by-id groups, duplicates adjacent.
    names = [ws.cell(row=r, column=name_col).value for r in range(2, 11)]
    assert names == [
        "GettyImages-0001.mov",
        "GettyImages-1073703562.mov",
        "GettyImages-2158057203.mov",
        "GettyImages-2158057203 (1).mxf",
        "GettyImages-2158057203 (2).mov",
        "GettyImages-356-30.mov",
        "GettyImages-356-31.mov",
        "GettyImages-9999.mov",
        "GettyImages-9999 (1).mxf",
    ]

    # Total Duration/Seconds only appear on the last row of each group.
    assert ws.cell(row=2, column=total_col).value == "00:00:00:20"
    assert ws.cell(row=2, column=seconds_col).value == 1

    assert ws.cell(row=3, column=total_col).value == "00:00:03:06"
    assert ws.cell(row=3, column=seconds_col).value == 4

    assert ws.cell(row=4, column=total_col).value is None  # mid-group
    assert ws.cell(row=5, column=total_col).value is None  # mid-group
    assert ws.cell(row=6, column=total_col).value == "00:00:05:03"  # 44+42+42 frames
    assert ws.cell(row=6, column=seconds_col).value == 6

    assert ws.cell(row=10, column=total_col).value == "00:00:03:00"  # 50+25 frames
    assert ws.cell(row=10, column=seconds_col).value == 3

    # Thick bottom border on the last row of each group, full row width.
    assert ws.cell(row=3, column=name_col).border.bottom.style == "thick"
    assert ws.cell(row=4, column=name_col).border.bottom.style != "thick"
    assert ws.cell(row=6, column=seconds_col).border.bottom.style == "thick"

    # Every row of a group is highlighted the same color, across the full
    # row width: GREEN if the group's total is 4s or less, YELLOW if over.
    # Group "2158057203" (rows 4-6) totals 6s -> yellow, on every row.
    assert ws.cell(row=4, column=name_col).fill.fgColor.rgb == "00FFFF00"
    assert ws.cell(row=5, column=name_col).fill.fgColor.rgb == "00FFFF00"
    assert ws.cell(row=6, column=seconds_col).fill.fgColor.rgb == "00FFFF00"
    # "GettyImages-1073703562.mov" totals exactly 4s -> not over 4s -> green.
    assert ws.cell(row=3, column=name_col).value == "GettyImages-1073703562.mov"
    assert ws.cell(row=3, column=name_col).fill.fgColor.rgb == "0092D050"
    assert ws.cell(row=3, column=1).fill.fgColor.rgb == "0092D050"  # full row width

    # 9 data rows (rows 2-10) -> row 11 is a blank spacer, summary starts row 12.
    assert ws.cell(row=11, column=name_col).value is None
    assert ws.cell(row=12, column=name_col).value == "Total clips"
    assert ws.cell(row=12, column=name_col + 1).value == 9
    assert ws.cell(row=13, column=name_col).value == "Total unique clips"
    assert ws.cell(row=13, column=name_col + 1).value == 6
    assert ws.cell(row=14, column=name_col).value == "Total clips >4s"
    assert ws.cell(row=14, column=name_col + 1).value == 1
    # Total Seconds / Total Seconds >4s are live formulas over the Seconds
    # column (rows 2-10), not hardcoded numbers, so edits to the sheet
    # recompute them.
    seconds_col_letter = get_column_letter(seconds_col)
    assert ws.cell(row=15, column=name_col).value == "Total Seconds"
    assert ws.cell(row=15, column=name_col + 1).value == f"=SUM({seconds_col_letter}2:{seconds_col_letter}10)"
    assert ws.cell(row=16, column=name_col).value == "Total Seconds for clips >4s"
    assert ws.cell(row=16, column=name_col + 1).value == f'=SUMIF({seconds_col_letter}2:{seconds_col_letter}10,">4")'

    assert results["Getty Videos"] == {
        "total_clips": 9,
        "total_unique_clips": 6,
        "total_clips_over_4s": 1,
        "total_seconds": 16,
        "total_seconds_over_4s": 6,
    }


def test_group_duplicates_leaves_other_sheets_untouched(sample_sorted_master, tmp_path):
    out_path = tmp_path / "out.xlsx"
    group_duplicates_workbook(str(sample_sorted_master), str(out_path))

    src = load_workbook(sample_sorted_master)
    out = load_workbook(out_path)

    for title in ["GFX", "Worksheet"]:
        src_ws, out_ws = src[title], out[title]
        assert out_ws.max_column == src_ws.max_column  # no Total Duration/Seconds added
        for r in range(1, src_ws.max_row + 1):
            for c in range(1, src_ws.max_column + 1):
                assert out_ws.cell(row=r, column=c).value == src_ws.cell(row=r, column=c).value


def test_group_duplicates_only_processes_default_sheets(sample_sorted_master, tmp_path):
    out_path = tmp_path / "out.xlsx"
    results = group_duplicates_workbook(str(sample_sorted_master), str(out_path))
    assert set(results.keys()) == {"Getty Videos", "AP"}


def test_group_duplicates_matches_sheet_names_case_insensitively(sample_sorted_master, tmp_path):
    # Real delivery files sometimes title this sheet "getty videos" (any
    # case) instead of our own "Getty Videos" convention.
    out_path = tmp_path / "out.xlsx"
    results = group_duplicates_workbook(str(sample_sorted_master), str(out_path), sheets=["ap", "GETTY VIDEOS"])
    assert set(results.keys()) == {"Getty Videos", "AP"}


def test_group_duplicates_matches_getty_pics_as_getty_stills_alias(tmp_path):
    # Real delivery files sometimes name the stills sheet "Getty pics"
    # instead of "Getty Stills" -- must be picked up by the default
    # sheet list without requiring --sheets.
    src_path = tmp_path / "src.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Getty pics"
    ws.append(["Clip Name", "Clip Duration"])
    ws.append(["GettyImages-1.jpg", "00:00:00:10"])
    ws.append(["GettyImages-1 (1).jpg", "00:00:00:10"])
    wb.save(src_path)

    out_path = tmp_path / "out.xlsx"
    results = group_duplicates_workbook(str(src_path), str(out_path))

    assert set(results.keys()) == {"Getty pics"}
    assert results["Getty pics"]["total_clips"] == 2
    assert results["Getty pics"]["total_unique_clips"] == 1


def test_group_duplicates_custom_sheets_override(sample_sorted_master, tmp_path):
    out_path = tmp_path / "out.xlsx"
    results = group_duplicates_workbook(str(sample_sorted_master), str(out_path), sheets=["GFX"])
    assert set(results.keys()) == {"GFX"}


def test_group_duplicates_falls_back_to_name_column(tmp_path):
    # Some real delivery files title the clip-name column "Name" instead of
    # "Clip Name" (the default). group_duplicates_workbook should still find
    # it without requiring name_column to be passed explicitly.
    src_path = tmp_path / "src.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "AP"
    ws.append(["Name", "Clip Duration"])
    ws.append(["apus_clip1.mxf", "00:00:00:10"])
    ws.append(["apus_clip1 (1).mxf", "00:00:00:10"])
    wb.save(src_path)

    out_path = tmp_path / "out.xlsx"
    results = group_duplicates_workbook(str(src_path), str(out_path), sheets=["AP"])

    assert results["AP"]["total_clips"] == 2
    assert results["AP"]["total_unique_clips"] == 1


def test_group_duplicates_falls_back_to_duration_column(tmp_path):
    # Some real delivery files title the duration column "Duration" instead
    # of "Clip Duration" (the default). group_duplicates_workbook should
    # still find it without requiring duration_column to be passed explicitly.
    src_path = tmp_path / "src.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "AP"
    ws.append(["Clip Name", "Duration"])
    ws.append(["apus_clip1.mxf", "00:00:00:10"])
    ws.append(["apus_clip1 (1).mxf", "00:00:00:10"])
    wb.save(src_path)

    out_path = tmp_path / "out.xlsx"
    results = group_duplicates_workbook(str(src_path), str(out_path), sheets=["AP"])

    assert results["AP"]["total_clips"] == 2
    assert results["AP"]["total_unique_clips"] == 1


def test_group_duplicates_collapses_trailing_number_tag_variant(tmp_path):
    # "GettyImages-627-122.mov" and "GettyImages-627-122.mov 25" are the
    # same clip -- a stray " 25" tag after the extension, not a different id.
    src_path = tmp_path / "src.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Getty Videos"
    ws.append(["Clip Name", "Clip Duration"])
    ws.append(["GettyImages-627-122.mov", "00:00:00:10"])
    ws.append(["GettyImages-627-122.mov 25", "00:00:00:10"])
    wb.save(src_path)

    out_path = tmp_path / "out.xlsx"
    results = group_duplicates_workbook(str(src_path), str(out_path), sheets=["Getty Videos"])

    assert results["Getty Videos"]["total_clips"] == 2
    assert results["Getty Videos"]["total_unique_clips"] == 1


def test_group_duplicates_collapses_codec_suffix_variant(tmp_path):
    # "GettyImages-1479649397.mov" and "GettyImages-1479649397_Apple_ProRes_422.mov"
    # are the same clip -- a codec tag baked in on re-export, not a different id.
    src_path = tmp_path / "src.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Getty Videos"
    ws.append(["Clip Name", "Clip Duration"])
    ws.append(["GettyImages-1479649397.mov", "00:00:00:10"])
    ws.append(["GettyImages-1479649397_Apple_ProRes_422.mov", "00:00:00:10"])
    wb.save(src_path)

    out_path = tmp_path / "out.xlsx"
    results = group_duplicates_workbook(str(src_path), str(out_path), sheets=["Getty Videos"])

    assert results["Getty Videos"]["total_clips"] == 2
    assert results["Getty Videos"]["total_unique_clips"] == 1


def test_group_duplicates_collapses_getty_upscale_suffix_variant(tmp_path):
    # Same Getty clip, exported with an optional "_S000" segment marker plus
    # an "_upscaleNN" tag -- not different clips.
    src_path = tmp_path / "src.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Getty Videos"
    ws.append(["Clip Name", "Clip Duration"])
    ws.append(["GettyImages-650878972_S000_upscale01.mov", "00:00:00:10"])
    ws.append(["GettyImages-650878972_S001_upscale02.mov", "00:00:00:10"])
    ws.append(["GettyImages-650878972_S002_upscale03.mov", "00:00:00:10"])
    ws.append(["GettyImages-650878972_upscale01.mov", "00:00:00:10"])
    wb.save(src_path)

    out_path = tmp_path / "out.xlsx"
    results = group_duplicates_workbook(str(src_path), str(out_path), sheets=["Getty Videos"])

    assert results["Getty Videos"]["total_clips"] == 4
    assert results["Getty Videos"]["total_unique_clips"] == 1


def test_group_duplicates_collapses_getty_render_prob4_comp_variant(tmp_path):
    # Getty QC/review-export junk ("Render", "Render <N>", "_prob4",
    # "_comp") re-appends the extension on every pass -- these are still
    # the same source clip and must group together, not split apart.
    src_path = tmp_path / "src.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Getty Videos"
    ws.append(["Clip Name", "Clip Duration"])
    ws.append(["GettyImages-1142748666.mov Render 1_prob4.mov", "00:00:00:10"])
    ws.append(["GettyImages-1142748666.mov Render_prob4.mov", "00:00:00:10"])
    ws.append(["GettyImages-2064403302_comp", "00:00:00:10"])
    ws.append(["GettyImages-2064403302_comp_2", "00:00:00:10"])
    wb.save(src_path)

    out_path = tmp_path / "out.xlsx"
    results = group_duplicates_workbook(str(src_path), str(out_path), sheets=["Getty Videos"])

    assert results["Getty Videos"]["total_clips"] == 4
    assert results["Getty Videos"]["total_unique_clips"] == 2


def test_group_duplicates_reuters_groups_by_tape_id(tmp_path):
    # Reuters rushes clips: same tape ("m<digits>" embedded in the name) but
    # different timecode ranges cut from it -- still the same source clip
    # for grouping purposes, unlike the default full-name dedup key.
    src_path = tmp_path / "src.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Reuters"
    ws.append(["Clip Name", "Clip Duration"])
    names = [
        "Clip from USA_ NEW YORK_ WORLD TRADE CENTER RUSHES COMPILATION TAPE ONE - 9_11 compilation_m909021_0-20-15-0-29-40.mov",
        "Clip from USA_ NEW YORK_ WORLD TRADE CENTER RUSHES COMPILATION TAPE ONE - 9_11 compilation_m909021_0-36-15-0-37-40.mov",
        "Clip from USA_ NEW YORK_ WORLD TRADE CENTER RUSHES COMPILATION TAPE ONE - 9_11 compilation_m909021_0-14-39-0-15-37 (1).mp4",
        "Clip from USA_ NEW YORK_ WORLD TRADE CENTER RUSHES COMPILATION TAPE THREE - 9_11 compilation_m917698_0-00-00-0-01-15_upscale.mp4",
        "Clip from USA_ NEW YORK_ WORLD TRADE CENTER RUSHES COMPILATION TAPE THREE - 9_11 compilation_m917698_0-03-15-0-03-35.mov",
        "Clip from USA_ NEW YORK_ WORLD TRADE CENTER RUSHES COMPILATION TAPE THREE - 9_11 compilation_m917698_0-08-00-0-08-18.mov",
        "Clip from USA_ NEW YORK_ WORLD TRADE CENTER RUSHES COMPILATION TAPE THREE - 9_11 compilation_m917698_0-32-45-0-33-05.mov",
        "Clip from USA_ NEW YORK_ WORLD TRADE CENTER RUSHES COMPILATION TAPE TWO - 9_11 compilation_m1241025_0-01-15-0-01-35.mov",
        "Clip from USA_ NEW YORK_ WORLD TRADE CENTER RUSHES COMPILATION TAPE TWO - 9_11 compilation_m1241025_0-35-40-0-36-00.mov",
    ]
    for name in names:
        ws.append([name, "00:00:00:10"])
    wb.save(src_path)

    out_path = tmp_path / "out.xlsx"
    results = group_duplicates_workbook(str(src_path), str(out_path), sheets=["Reuters"])

    assert results["Reuters"]["total_clips"] == 9
    assert results["Reuters"]["total_unique_clips"] == 3  # tapes m909021, m917698, m1241025


def test_group_duplicates_blank_duration_counts_as_zero(tmp_path):
    # Real delivery files sometimes leave Clip Duration blank on a row.
    # It should count as 0 frames, not crash the whole batch.
    src_path = tmp_path / "src.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "AP"
    ws.append(["Clip Name", "Clip Duration"])
    ws.append(["apus_clip1.mxf", "00:00:00:10"])
    ws.append(["apus_clip1 (1).mxf", None])
    wb.save(src_path)

    out_path = tmp_path / "out.xlsx"
    results = group_duplicates_workbook(str(src_path), str(out_path), sheets=["AP"])

    assert results["AP"]["total_clips"] == 2
    assert results["AP"]["total_unique_clips"] == 1
    assert results["AP"]["total_seconds"] == 1  # 10 frames @25fps -> ceil to 1s


def test_group_output_is_always_styled(sample_sorted_master, tmp_path):
    # Styling is unconditional now: every processed and pass-through sheet
    # (all but the verbatim "Worksheet" backup) gets black/white bold
    # headers and header-keyed column widths, no flags required.
    out = tmp_path / "out.xlsx"
    group_duplicates_workbook(str(sample_sorted_master), str(out))
    wb = load_workbook(str(out))
    styled = [t for t in wb.sheetnames if t != "Worksheet"]
    assert styled
    for t in styled:
        ws = wb[t]
        name_c = next(c for c in range(1, ws.max_column + 1)
                      if ws.cell(row=1, column=c).value == "Clip Name")
        h = ws.cell(row=1, column=name_c)
        assert h.fill.fgColor.rgb == "FF000000"
        assert h.font.bold and h.font.color.rgb == "FFFFFFFF"
        assert 10 <= ws.column_dimensions[get_column_letter(name_c)].width <= 60


def test_group_styling_leaves_worksheet_backup_alone(sample_sorted_master, tmp_path):
    out_path = tmp_path / "out.xlsx"
    group_duplicates_workbook(str(sample_sorted_master), str(out_path))

    src = load_workbook(sample_sorted_master)["Worksheet"]
    backup = load_workbook(out_path)["Worksheet"]
    assert backup.column_dimensions["B"].width == src.column_dimensions["B"].width
