import io

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from web.auth import hash_password
from web.main import app


def _logged_in_client(monkeypatch):
    monkeypatch.setenv("BIGMEDIA_WEB_PASSWORD_HASH", hash_password("pw"))
    client = TestClient(app)
    client.post("/login", data={"password": "pw"})
    return client


def _xlsx_bytes(names, header="Clip Name", extra_sheets=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Master"
    ws.append([header])
    for name in names:
        ws.append([name])
    for title, rows in (extra_sheets or {}).items():
        s = wb.create_sheet(title=title)
        for row in rows:
            s.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_analyze_returns_category_annotations(monkeypatch):
    client = _logged_in_client(monkeypatch)
    content = _xlsx_bytes(["BM1234_x.mxf", "BM5678_y.mxf", "random_file_xyz.mov"])
    files = {"files": ("master.xlsx", content, "application/octet-stream")}
    response = client.post("/commands/sort/analyze", data={"name_column": "Clip Name"}, files=files)
    assert response.status_code == 200
    body = response.json()
    assert body["annotations"]["categories"]["AP"] == 2
    assert body["annotations"]["categories"]["3rd parties"] == 1
    assert "AP" in body["suggestions"]["categories"]
    assert "Reuters" not in body["suggestions"]["categories"]


def test_analyze_returns_headers_and_sheets(monkeypatch):
    client = _logged_in_client(monkeypatch)
    content = _xlsx_bytes(["x.mov"], extra_sheets={"Getty Videos": [["Clip Name"], ["g.mov"]]})
    files = {"files": ("master.xlsx", content, "application/octet-stream")}
    response = client.post("/commands/sort/analyze", data={}, files=files)
    body = response.json()
    assert body["headers"] == ["Clip Name"]
    names = {s["name"]: s["rows"] for s in body["sheets"]}
    assert names == {"Master": 1, "Getty Videos": 1}


def test_analyze_union_across_multiple_files(monkeypatch):
    client = _logged_in_client(monkeypatch)
    a = _xlsx_bytes(["BM1_x.mxf"], header="Clip Name")
    b = _xlsx_bytes(["BM2_y.mxf"], header="Name")
    files = [
        ("files", ("a.xlsx", a, "application/octet-stream")),
        ("files", ("b.xlsx", b, "application/octet-stream")),
    ]
    response = client.post("/commands/sort/analyze", data={"name_column": "Clip Name"}, files=files)
    body = response.json()
    assert set(body["headers"]) == {"Clip Name", "Name"}
    assert body["annotations"]["categories"]["AP"] == 2
    assert any("differ" in w for w in body["warnings"])


def test_analyze_missing_name_column_warns_but_200(monkeypatch):
    client = _logged_in_client(monkeypatch)
    content = _xlsx_bytes(["BM1234_x.mxf"], header="Media File")
    files = {"files": ("master.xlsx", content, "application/octet-stream")}
    response = client.post("/commands/sort/analyze", data={"name_column": "Clip Name"}, files=files)
    assert response.status_code == 200
    assert response.json()["warnings"]


def test_analyze_generic_command_has_no_suggestions(monkeypatch):
    client = _logged_in_client(monkeypatch)
    content = _xlsx_bytes(["x.mov"])
    files = {"files": ("master.xlsx", content, "application/octet-stream")}
    response = client.post("/commands/dedupe/analyze", data={}, files=files)
    body = response.json()
    assert body["headers"] == ["Clip Name"]
    assert body["suggestions"] == {}


def test_analyze_unknown_slug_404(monkeypatch):
    client = _logged_in_client(monkeypatch)
    files = {"files": ("m.xlsx", _xlsx_bytes(["x"]), "application/octet-stream")}
    response = client.post("/commands/nope/analyze", data={}, files=files)
    assert response.status_code == 404


def test_analyze_rejects_non_xlsx(monkeypatch):
    client = _logged_in_client(monkeypatch)
    files = {"files": ("notes.txt", b"hello", "text/plain")}
    response = client.post("/commands/sort/analyze", data={"name_column": "Clip Name"}, files=files)
    assert response.status_code == 400
    assert "notes.txt" in response.json()["error"]


def test_analyze_requires_login(monkeypatch):
    monkeypatch.setenv("BIGMEDIA_WEB_PASSWORD_HASH", hash_password("pw"))
    client = TestClient(app, follow_redirects=False)
    files = {"files": ("master.xlsx", _xlsx_bytes(["x.mov"]), "application/octet-stream")}
    response = client.post("/commands/sort/analyze", data={"name_column": "Clip Name"}, files=files)
    assert response.status_code == 303


def test_sort_run_accepts_multiple_categories_checkboxes(monkeypatch):
    client = _logged_in_client(monkeypatch)
    content = _xlsx_bytes([
        "BM1234_something.mxf",                  # AP
        "045_SO_EP18_01_3DExplainer_TXLS.mov",   # GFX
        "shutterstock_777.mp4",                  # Shutterstock
    ])
    files = {"files": ("master.xlsx", content, "application/octet-stream")}
    response = client.post(
        "/commands/sort",
        data={"name_column": "Clip Name", "categories": ["AP", "GFX"]},
        files=files,
    )
    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.content))
    assert "AP" in wb.sheetnames
    assert "GFX" in wb.sheetnames
    assert "Shutterstock" not in wb.sheetnames
    assert "3rd parties" in wb.sheetnames
