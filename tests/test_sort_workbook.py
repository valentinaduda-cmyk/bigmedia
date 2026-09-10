import pytest
from openpyxl import Workbook, load_workbook
from bigmedia.classify import CATEGORY_ORDER
from bigmedia.sort_workbook import count_categories, sort_workbook
from bigmedia.xlsx_utils import find_column


def test_count_categories_tallies_clips_per_category(sample_master):
    counts = count_categories([str(sample_master)])

    assert counts["AP"] == 2
    assert counts["GFX"] == 1
    assert counts["Getty Unknown"] == 1
    assert counts["Getty Videos"] == 1
    assert counts["Shutterstock"] == 1
    assert counts["3rd parties"] == 1
    # Every known category is present, zero when absent from the file.
    assert counts["Reuters"] == 0
    assert set(counts.keys()) == set(CATEGORY_ORDER)


def test_count_categories_aggregates_across_multiple_files(sample_master):
    counts = count_categories([str(sample_master), str(sample_master)])

    assert counts["AP"] == 4
    assert counts["3rd parties"] == 2


def test_sort_workbook_creates_backup_and_category_sheets(sample_master, tmp_path):
    out_path = tmp_path / "out.xlsx"
    counts = sort_workbook(str(sample_master), str(out_path))

    wb = load_workbook(out_path)

    # Worksheet backup must be first and match the source exactly, including
    # "Source Reel Name" — sort no longer touches that column.
    assert wb.sheetnames[0] == "Worksheet"
    src = load_workbook(sample_master).active
    backup = wb["Worksheet"]
    for r in range(1, src.max_row + 1):
        for c in range(1, src.max_column + 1):
            assert backup.cell(row=r, column=c).value == src.cell(row=r, column=c).value

    assert counts["AP"] == 2  # includes the "(1)" duplicate — sort doesn't dedupe
    assert counts["GFX"] == 1
    assert counts["Getty Unknown"] == 1
    assert counts["Getty Videos"] == 1
    assert counts["Shutterstock"] == 1
    assert counts["3rd parties"] == 1


def test_sort_workbook_leaves_source_reel_name_untouched(sample_master, tmp_path):
    out_path = tmp_path / "out.xlsx"
    sort_workbook(str(sample_master), str(out_path))

    wb = load_workbook(out_path)
    src = load_workbook(sample_master).active
    source_col_idx = find_column(src, "Source Reel Name")
    name_col_idx = find_column(src, "Clip Name")

    for cat in ["AP", "GFX", "Getty Unknown", "Getty Videos", "Shutterstock", "3rd parties"]:
        ws = wb[cat]
        for r in range(2, ws.max_row + 1):
            name = ws.cell(row=r, column=name_col_idx).value
            for src_r in range(2, src.max_row + 1):
                if src.cell(row=src_r, column=name_col_idx).value == name:
                    assert ws.cell(row=r, column=source_col_idx).value == src.cell(row=src_r, column=source_col_idx).value
                    break


def test_sort_workbook_custom_name_column(sample_master, tmp_path):
    out_path = tmp_path / "out.xlsx"
    counts = sort_workbook(str(sample_master), str(out_path), name_column="Clip Name")
    assert sum(counts.values()) == 7


def test_sort_workbook_falls_back_to_name_column(tmp_path):
    # Some real delivery files title the clip-name column "Name" instead of
    # "Clip Name" (the default). sort_workbook should still find it without
    # requiring name_column to be passed explicitly.
    src_path = tmp_path / "src.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Name", "Source Reel Name"])
    ws.append(["apus_clip1.mxf", ""])
    ws.append(["shutterstock_777.mp4", ""])
    wb.save(src_path)

    out_path = tmp_path / "out.xlsx"
    counts = sort_workbook(str(src_path), str(out_path))
    assert counts["AP"] == 1
    assert counts["Shutterstock"] == 1


def test_sort_workbook_skip_categories_omits_sheet_and_falls_back(sample_master, tmp_path):
    # A category turned off keeps its classification rules — its clips just
    # land on "3rd parties" for manual review instead of getting their own
    # sheet. Nothing is dropped.
    out_path = tmp_path / "out.xlsx"
    counts = sort_workbook(str(sample_master), str(out_path), skip_categories=["Shutterstock"])

    wb = load_workbook(out_path)
    assert "Shutterstock" not in wb.sheetnames
    assert "3rd parties" in counts
    assert "Shutterstock" not in counts
    assert counts["3rd parties"] == 2  # its own row + the shutterstock one
    assert counts["AP"] == 2
    assert sum(counts.values()) == 7  # every non-empty row still accounted for


def test_sort_workbook_categories_allowlist(sample_master, tmp_path):
    out_path = tmp_path / "out.xlsx"
    counts = sort_workbook(str(sample_master), str(out_path), categories=["AP", "GFX"])

    wb = load_workbook(out_path)
    assert wb.sheetnames == ["Worksheet", "AP", "GFX", "3rd parties"]
    assert counts["AP"] == 2
    assert counts["GFX"] == 1
    # Getty Unknown + Getty Videos + Shutterstock + the already-unclassified row
    assert counts["3rd parties"] == 4


def test_sort_workbook_category_names_are_case_insensitive(sample_master, tmp_path):
    out_path = tmp_path / "out.xlsx"
    counts = sort_workbook(str(sample_master), str(out_path), skip_categories=["shutterstock", "GETTY VIDEOS"])
    assert "Shutterstock" not in counts
    assert "Getty Videos" not in counts
    assert counts["3rd parties"] == 3


def test_sort_workbook_rejects_both_selection_flags(sample_master, tmp_path):
    with pytest.raises(ValueError):
        sort_workbook(str(sample_master), str(tmp_path / "out.xlsx"),
                      categories=["AP"], skip_categories=["GFX"])


def test_sort_workbook_rejects_unknown_category_name(sample_master, tmp_path):
    # A typo must fail loudly rather than silently sorting nothing out.
    with pytest.raises(ValueError) as exc:
        sort_workbook(str(sample_master), str(tmp_path / "out.xlsx"), skip_categories=["Fox News"])
    assert "Fox News" in str(exc.value)


def test_sort_workbook_cannot_skip_the_fallback_category(sample_master, tmp_path):
    # "3rd parties" is where every turned-off category's clips land, so it
    # can't itself be turned off.
    with pytest.raises(ValueError):
        sort_workbook(str(sample_master), str(tmp_path / "out.xlsx"), skip_categories=["3rd parties"])


def test_sort_output_sheets_are_always_styled(sample_master, tmp_path):
    from openpyxl import load_workbook
    out = tmp_path / "out.xlsx"
    sort_workbook(str(sample_master), str(out))
    wb = load_workbook(str(out))
    cat_sheets = [t for t in wb.sheetnames if t != "Worksheet"]
    assert cat_sheets
    ref = None
    for t in cat_sheets:
        ws = wb[t]
        # header row: black fill, white bold
        h = ws.cell(row=1, column=1)
        assert h.fill.fgColor.rgb == "FF000000"
        assert h.font.bold and h.font.color.rgb == "FFFFFFFF"
        # widths in range and shared across sheets
        w = ws.column_dimensions["A"].width
        assert 10 <= w <= 60
        ref = ref or w
        assert ws.column_dimensions["A"].width == ref
        # uniform data font (row 2+ all one name), when the sheet has data
        fonts = {ws.cell(row=r, column=1).font.name
                 for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value}
        assert len(fonts) <= 1


def test_sort_worksheet_backup_stays_verbatim(sample_master, tmp_path):
    from openpyxl import load_workbook
    out = tmp_path / "out.xlsx"
    sort_workbook(str(sample_master), str(out))
    wb = load_workbook(str(out))
    src = load_workbook(str(sample_master)).active
    bak = wb["Worksheet"]
    # header fill NOT forced to black on the backup
    assert bak.cell(row=1, column=1).fill.fgColor.rgb != "FF000000" or \
           src.cell(row=1, column=1).fill.fgColor.rgb == "FF000000"
    # values identical
    for r in range(1, src.max_row + 1):
        for c in range(1, src.max_column + 1):
            assert bak.cell(row=r, column=c).value == src.cell(row=r, column=c).value


def test_analyze_sort_suggests_only_nonzero_categories(tmp_path):
    from bigmedia.sort_workbook import analyze_sort
    wb = Workbook()
    ws = wb.active
    ws.append(["Clip Name"])
    ws.append(["BM1234_x.mxf"])       # AP
    ws.append(["shutterstock_9.mp4"]) # Shutterstock
    path = tmp_path / "m.xlsx"
    wb.save(path)

    result = analyze_sort([str(path)])
    assert set(result["suggestions"]["categories"]) == {"AP", "Shutterstock"}
    assert result["annotations"]["categories"]["AP"] == 1
    assert result["annotations"]["categories"]["Reuters"] == 0
    assert "warnings" not in result


def test_analyze_sort_missing_name_column_warns_not_raises(tmp_path):
    from bigmedia.sort_workbook import analyze_sort
    wb = Workbook()
    wb.active.append(["Something Else"])
    path = tmp_path / "m.xlsx"
    wb.save(path)

    result = analyze_sort([str(path)], name_column="Clip Name")
    assert result["warnings"]
    assert "suggestions" not in result
