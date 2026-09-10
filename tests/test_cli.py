import shutil

import pytest
from openpyxl import load_workbook

from bigmedia.cli import main, _iter_xlsx_inputs


def test_iter_xlsx_inputs_skips_excel_lock_files(tmp_path):
    (tmp_path / "a.xlsx").touch()
    (tmp_path / "b.xlsx").touch()
    (tmp_path / "~$a.xlsx").touch()  # Excel's lock file for an open "a.xlsx"
    (tmp_path / "notes.txt").touch()

    found = [f.name for f in _iter_xlsx_inputs(tmp_path)]
    assert found == ["a.xlsx", "b.xlsx"]


def test_iter_xlsx_inputs_single_file_passthrough(sample_master):
    assert _iter_xlsx_inputs(sample_master) == [sample_master]


def test_cmd_sort_batch_processes_every_file_in_folder(sample_master, tmp_path):
    in_dir = tmp_path / "in"
    in_dir.mkdir()
    shutil.copy(sample_master, in_dir / "EP1_master.xlsx")
    shutil.copy(sample_master, in_dir / "EP2_master.xlsx")

    main(["sort", str(in_dir)])

    out1 = in_dir / "EP1_master_sorted.xlsx"
    out2 = in_dir / "EP2_master_sorted.xlsx"
    assert out1.exists()
    assert out2.exists()
    assert load_workbook(out1)["AP"].max_row == 3  # header + 2 AP rows


def test_cmd_sort_batch_with_output_folder(sample_master, tmp_path):
    in_dir = tmp_path / "in"
    in_dir.mkdir()
    shutil.copy(sample_master, in_dir / "EP1_master.xlsx")

    out_dir = tmp_path / "out"
    main(["sort", str(in_dir), "-o", str(out_dir)])

    assert (out_dir / "EP1_master_sorted.xlsx").exists()
    assert not (in_dir / "EP1_master_sorted.xlsx").exists()


def test_cmd_dedupe_batch_processes_every_file_in_folder(sample_master, tmp_path):
    in_dir = tmp_path / "in"
    in_dir.mkdir()
    shutil.copy(sample_master, in_dir / "EP1_master.xlsx")

    main(["dedupe", str(in_dir)])

    assert (in_dir / "EP1_master_deduped.xlsx").exists()


def test_cmd_fu_grid_batch_processes_every_file_in_folder(sample_third_parties, tmp_path):
    in_dir = tmp_path / "in"
    in_dir.mkdir()
    shutil.copy(sample_third_parties, in_dir / "EP1_sorted.xlsx")

    main(["fu-grid", str(in_dir)])

    out = in_dir / "EP1_sorted_fu_grid.xlsx"
    assert out.exists()
    assert load_workbook(out).sheetnames == ["3rd parties", "FU grid"]


def test_cmd_sort_skip_categories_flag(sample_master, tmp_path):
    out = tmp_path / "out.xlsx"
    main(["sort", str(sample_master), "-o", str(out), "--skip-categories", "FOX, Shutterstock"])

    wb = load_workbook(out)
    assert "Shutterstock" not in wb.sheetnames
    assert "FOX" not in wb.sheetnames
    assert wb["3rd parties"].max_row == 3  # header + own row + the shutterstock one


def test_cmd_sort_categories_flag(sample_master, tmp_path):
    out = tmp_path / "out.xlsx"
    main(["sort", str(sample_master), "-o", str(out), "--categories", "AP,GFX"])

    assert load_workbook(out).sheetnames == ["Worksheet", "AP", "GFX", "3rd parties"]


def test_removed_styling_flags_no_longer_parse(sample_master, tmp_path):
    # The old opt-in styling flags were dropped, not kept as no-ops --
    # passing one is an argparse error.
    out = tmp_path / "out.xlsx"
    with pytest.raises(SystemExit):
        main(["sort", str(sample_master), "-o", str(out), "--autofit"])
    with pytest.raises(SystemExit):
        main(["group", str(sample_master), "-o", str(out), "--uniform-header"])
